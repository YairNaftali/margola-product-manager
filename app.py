import csv, html, io, json, os, re, uuid, time, mimetypes, urllib.parse, urllib.request, ssl, certifi
from datetime import datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import urlparse, parse_qs
import openpyxl

APP_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(APP_DIR, "data")
UPLOAD_DIR = os.path.join(APP_DIR, "uploads")
PRODUCTS_PATH = os.path.join(DATA_DIR, "products.json")
os.makedirs(DATA_DIR, exist_ok=True)
os.makedirs(UPLOAD_DIR, exist_ok=True)

def clean(v):
    if v is None: return ""
    if isinstance(v, float) and v.is_integer(): return str(int(v))
    return str(v).strip()

def slugify(v):
    v = clean(v).lower()
    v = re.sub(r"[^a-z0-9]+", "-", v)
    return re.sub(r"-+", "-", v).strip("-")

ALT_TEXT_PROTECTED_TERMS = {"AB": "AB", "DK": "Dark", "2XAB": "2X AB"}

def make_alt_text(title):
    t = clean(title)
    t = re.sub(r"\s*[–-]\s*\d+\s*Gross.*$", "", t, flags=re.I)
    t = re.sub(r"\s*\(\d[\d,]*\s*pcs\)\s*$", "", t, flags=re.I)
    t = re.sub(r"\s+", " ", t).strip(" -–")
    if not t.isupper():
        return t
    t = re.sub(r"\bLT\.?\s*", "LIGHT ", t)
    t = re.sub(r"\s+", " ", t).strip()
    def repl(m):
        w = m.group(0)
        key = w.rstrip(".").upper()
        return ALT_TEXT_PROTECTED_TERMS.get(key, w.title())
    return re.sub(r"[A-Za-z0-9/.'-]+", repl, t)

def strip_html(raw):
    txt = re.sub(r"<[^>]+>", " ", raw or "")
    txt = html.unescape(txt)
    txt = re.sub(r"\s+", " ", txt).strip()
    return re.sub(r"\s+([.,;:!?])", r"\1", txt)

def fix_shouty_caps(text):
    """Title-cases any run of 3+ consecutive ALL-CAPS words (several
    description templates -- Roller Beads, 2/3 Cut -- have ALL-CAPS headings
    baked into the stored HTML itself, not just CSS text-transform)."""
    def repl(m):
        words = m.group(0).split()
        out = []
        for w in words:
            key = w.rstrip(".,").upper()
            out.append(ALT_TEXT_PROTECTED_TERMS.get(key, w.title()))
        return " ".join(out)
    return re.sub(r"\b(?:[A-Z][A-Z'/-]*\s+){2,}[A-Z][A-Z'/-]*\b", repl, text)

def truncate_words(txt, limit=160):
    """Cut at the last full sentence inside the limit if there is one far
    enough in to be worthwhile; otherwise cut at the last full word. Never
    cut mid-word -- this was wrong in an early version of this script and
    produced descriptions ending on a dangling half-word."""
    if len(txt) <= limit:
        return txt
    window = txt[:limit]
    last_period = window.rfind(". ")
    if last_period > 80:
        return window[:last_period + 1]
    return window.rsplit(" ", 1)[0].rstrip(",.;:- ")

def gen_seo_title(title):
    seo_title = make_alt_text(title)
    if seo_title == title:
        # Shopify silently stores null if seo.title is byte-identical to the
        # product's own title (userErrors: [] but the write no-ops) -- confirmed
        # live 2026-07-23, hit 204/393 products in one run. See SEO_TITLE_DESCRIPTION_SPEC.md.
        seo_title = f"{title} | Margola"
    return seo_title

def _build_seo_description(title, description_html):
    ct = make_alt_text(title)
    raw_plain = strip_html(description_html)
    if not raw_plain:
        text = f"Shop {ct} from Margola — quality Czech glass beads & findings for jewelry, costume, and craft projects."
        return text, {"used_fallback": True, "shouty_caps_fixed": False, "awkward_truncation": False}
    fixed = fix_shouty_caps(raw_plain)
    shouty_changed = fixed != raw_plain
    if fixed.lower().startswith(ct.lower()):
        fixed = fixed[len(ct):].strip(" -–")
    truncated = truncate_words(fixed, 160)
    awkward = len(truncated) == 160 and truncated[-1:] not in (".", "!", "?")
    return truncated, {"used_fallback": False, "shouty_caps_fixed": shouty_changed, "awkward_truncation": awkward}

def gen_description(title, description_html):
    return _build_seo_description(title, description_html)[0]

def shopify_list_products_for_seo():
    nodes, cursor = [], None
    while True:
        q = """query($after:String){ products(first:100, after:$after){ edges{ cursor node{ id handle title descriptionHtml seo{title description} } } pageInfo{ hasNextPage } } }"""
        data = shopify_graphql(q, {"after": cursor})
        edges = data["products"]["edges"]
        nodes.extend(e["node"] for e in edges)
        if not data["products"]["pageInfo"]["hasNextPage"] or not edges: break
        cursor = edges[-1]["cursor"]
    return nodes

def shopify_seo_proposals():
    # Reads the whole live catalog directly from Shopify (not local products.json,
    # which only holds the currently-loaded review batch) -- per the spec, "already
    # has a value" isn't proof it's still correct (see the 7/9 Rhinestone incident),
    # so this needs to be safe and normal to re-run against every product, not just
    # newly imported ones.
    proposals = []
    for node in shopify_list_products_for_seo():
        title = node.get("title") or ""
        description_html = node.get("descriptionHtml") or ""
        current = node.get("seo") or {}
        seo_title = gen_seo_title(title)
        seo_description, desc_flags = _build_seo_description(title, description_html)
        flags = []
        if seo_title == title:
            flags.append("seo_title_still_matches_title")
        if desc_flags["awkward_truncation"]:
            flags.append("awkward_truncation")
        if desc_flags["shouty_caps_fixed"]:
            flags.append("shouty_caps_fixed")
        if desc_flags["used_fallback"]:
            flags.append("no_description_used_fallback")
        current_title, current_description = current.get("title") or "", current.get("description") or ""
        changed = seo_title != current_title or seo_description != current_description
        proposals.append({
            "id": node["id"], "handle": node["handle"], "title": title,
            "current_seo_title": current_title, "current_seo_description": current_description,
            "seo_title": seo_title, "seo_description": seo_description,
            "flags": flags, "changed": changed,
        })
    return proposals

def shopify_apply_seo(items):
    applied, errors = [], []
    for item in items:
        handle = clean(item.get("handle")); product_id = clean(item.get("id"))
        seo_title, seo_description = item.get("seo_title", ""), item.get("seo_description", "")
        if not product_id:
            errors.append({"handle": handle, "error": "Missing product id"}); continue
        try:
            m = """mutation($product:ProductUpdateInput!){ productUpdate(product:$product){ product{ id } userErrors{ field message } } }"""
            d = shopify_graphql(m, {"product": {"id": product_id, "seo": {"title": seo_title, "description": seo_description}}})["productUpdate"]
            if d.get("userErrors"): raise RuntimeError(json.dumps(d["userErrors"]))
            applied.append({"handle": handle, "seo_title": seo_title, "seo_description": seo_description})
        except Exception as e:
            errors.append({"handle": handle, "error": str(e)})
    return {"applied": applied, "errors": errors}

def normalize_header(h):
    return re.sub(r"[^a-z0-9]+", "_", clean(h).lower()).strip("_")

def money(v):
    if v in (None, ""): return ""
    try: return f"{float(v):.2f}"
    except Exception: return ""

def weight_oz(v):
    text = clean(v).lower().replace("ounces", "oz").replace("ounce", "oz")
    m = re.search(r"[-+]?\d*\.?\d+", text)
    return m.group(0) if m else ""

def get_first(row, names):
    for n in names:
        k = normalize_header(n)
        if k in row and row[k] not in (None, ""):
            return row[k]
    return ""

def row_dict(headers, values):
    return {normalize_header(headers[i]): values[i] if i < len(values) else None for i in range(len(headers))}

def spreadsheet_types_path(): return os.path.join(DATA_DIR, "spreadsheet_types.json")

_SPREADSHEET_TYPES_CACHE = None
def load_spreadsheet_types(force=False):
    global _SPREADSHEET_TYPES_CACHE
    if force or _SPREADSHEET_TYPES_CACHE is None:
        path = spreadsheet_types_path()
        _SPREADSHEET_TYPES_CACHE = json.load(open(path, encoding="utf-8")) if os.path.exists(path) else []
    return _SPREADSHEET_TYPES_CACHE

def spreadsheet_type_by_id(type_id):
    return next((t for t in load_spreadsheet_types() if t["id"] == type_id), None)

def _infer_fields(t):
    # Strips the matching/sync keys, leaving just what infer() has always returned.
    return {k: t[k] for k in ("collection","title_prefix","title_suffix","subcategory","bead_shape","type","color_type","factory_qty_standard","mini_qty_standard","populate_bead_shape_from_shape","shopify_category","vendor")}

def infer(sheet_name, color_name, source_filename="", forced_type_id=None):
    s = sheet_name.lower()
    s_compact = re.sub(r"[^a-z0-9]", "", s)
    fn = source_filename.lower()
    c = clean(color_name).lower()
    types = load_spreadsheet_types()

    if forced_type_id:
        t = spreadsheet_type_by_id(forced_type_id)
        if not t:
            raise ValueError(f"Unknown spreadsheet type id {forced_type_id!r}. Known: {[x['id'] for x in types]}")
    else:
        t = None
        for cand in types:
            if any(m in s_compact for m in cand.get("sheet_match", [])) or any(m in fn for m in cand.get("filename_match", [])):
                t = cand; break

    if not t:
        return {"type_id":"","collection":"","title_prefix":"","title_suffix":"","subcategory":"","bead_shape":"","type":"","color_type":"","factory_qty_standard":"","mini_qty_standard":"","populate_bead_shape_from_shape":False,"shopify_category":"","vendor":""}

    fields = _infer_fields(t)
    if t["id"] == "roller-beads":
        # Depends on scanning both sheet name and color name together -- not a
        # fit for a static JSON field, kept as the one piece of special-cased logic.
        if "transparent" in s or "transparent" in c or "transpaent" in c: fields["color_type"] = "Transparent"
        elif "opaque" in s or "opaque" in c: fields["color_type"] = "Opaque"
    elif t["id"] == "cameos-intaglios":
        # Cameo vs Intaglio isn't its own sheet column -- it's a word embedded in
        # the raw Color Name text (e.g. "Intaglio Crystal Single Rose" vs "Cameo
        # Black"), confirmed 2026-08-03 against every live product's title. Default
        # from the JSON is "Cameo"; override to "Intaglio" when the color name says so.
        if "intaglio" in c: fields["type"] = "Intaglio"
    fields["type_id"] = t["id"]
    return fields

def title_for(prefix, descriptor, size, color_name):
    # No assumptions: do not invent manufacturer/brand.
    return " ".join([x for x in [clean(prefix), clean(descriptor), clean(size), clean(color_name)] if x])

def image_for(factory_style, color_name, subcategory, size):
    base = clean(factory_style) or " ".join([subcategory, size, color_name])
    return f"{slugify(base)}.jpg" if base else ""


def build_variants(factory_style, factory_price, factory_weight, factory_qty_desc, mini_style, mini_price, mini_weight, mini_qty):
    variants = []

    # No assumptions:
    # A pack variant only exists if it has its own SKU AND at least some real pack data.
    # Do not copy Mini into Factory or Factory into Mini.
    if clean(mini_style) and (clean(mini_price) or clean(mini_weight) or clean(mini_qty)):
        variants.append({
            "name": "Mini Pack",
            "option_value": "mini-pack",
            "sku": clean(mini_style),
            "price": money(mini_price),
            "weight_oz": weight_oz(mini_weight),
            "quantity": clean(mini_qty),
        })

    if clean(factory_style) and (clean(factory_price) or clean(factory_weight) or clean(factory_qty_desc)):
        variants.append({
            "name": "Factory Pack",
            "option_value": "factory-pack",
            "sku": clean(factory_style),
            "price": money(factory_price),
            "weight_oz": weight_oz(factory_weight),
            "quantity": clean(factory_qty_desc),
        })

    return variants

def validate(p):
    variants = p.get("variants", [])

    checks = {
        "Title": bool(p.get("title")),
        "Handle": bool(p.get("handle")),
        "At least one variant": bool(variants),
        "All variant SKUs": bool(variants) and all(v.get("sku") for v in variants),
        "All variant prices": bool(variants) and all(v.get("price") for v in variants),
        "All variant weights": bool(variants) and all(v.get("weight_oz") for v in variants),
        "Collection": bool(p.get("collection")),
        "Subcategory": bool(p.get("subcategory")),
        "Color type": bool(p.get("color_type")),
        "Bead shape": bool(p.get("bead_shape")),
        "Image filename": bool(p.get("image_filename")),
        "Description": bool(p.get("generated_description") or p.get("source_description")),
    }

    score = int(round(sum(1 for ok in checks.values() if ok) / len(checks) * 100))
    warnings = [k for k, ok in checks.items() if not ok]

    if p.get("notes"):
        warnings.append(p["notes"])

    return {"checks": checks, "score": score, "warnings": warnings}

def parse_xlsx(path, forced_type_id=None):
    wb = openpyxl.load_workbook(path, data_only=True)
    products = []
    for ws in wb.worksheets:
        row1 = [c.value for c in ws[1]]
        if not any(row1): continue
        row2 = [c.value for c in ws[2]] if ws.max_row >= 2 else []
        # Some sheets (e.g. 2 Cut / seed / bugle bead templates) use a merged two-row
        # header: row 1 has group labels (e.g. "Factory Pack Style Number"), row 2 has
        # the specific field name for grouped columns (e.g. "Color Number"). Detect that
        # by row 2 containing recognizable field-name text, since a genuine first data
        # row would never literally contain the words "Color Number"/"Color Name".
        row2_norm = {normalize_header(v) for v in row2 if v}
        two_row_header = "color_number" in row2_norm or "color_name" in row2_norm
        if two_row_header:
            width = max(len(row1), len(row2))
            row1 += [None] * (width - len(row1))
            row2 += [None] * (width - len(row2))
            headers = [row2[i] if row2[i] not in (None, "") else row1[i] for i in range(width)]
            data_start = 3
        else:
            headers = row1
            data_start = 2
        for row_num, vals in enumerate(ws.iter_rows(min_row=data_start, values_only=True), start=data_start):
            if not any(vals): continue
            row = row_dict(headers, vals)
            factory_style = get_first(row, ["FACTORY PACK STYLE #","FACTORY PACK       STYLE #","FACTORY PACK STYLE","FACTORY PACK STYLE NUMBER"])
            if not clean(factory_style): continue
            size = get_first(row, ["BEAD SIZE (FILTER)","BEAD SIZE FILTER","STONE SIZE","BEAD SIZE DIAMETER  MM","STONE SIZE DIAMETER MM","BEAD SIZE","SIZE","Milimeter size"])
            size_mm = get_first(row, ["BEAD SIZE DIAMETER  MM","BEAD SIZE DIAMETER MM","STONE SIZE DIAMETER MM"])
            color_number = get_first(row, ["COLOR NUMBER"])
            color_name = get_first(row, ["COLOR NAME"])
            color_type_explicit = get_first(row, ["COLOR TYPE (IS A FILTER ON THE WEBITE)","COLOR TYPE (IS A FILTER ON THE WEBSITE)","COLOR TYPE"])
            shape = get_first(row, ["SHAPE"])
            source_description = get_first(row, ["DESCRIPTION"])
            factory_qty = get_first(row, ["FACTORY PACK UNIT QUANTITY"])
            factory_qty_desc = get_first(row, ["UNIT QUANTITY DESCRIPTION","UNIT QUANTITY DESCRIPTION FACTORY PACK"])
            factory_price = get_first(row, ["UNIT PRICE PER FACTORY PACK","UNIT PRICE PER BAG","FACTORY PACK PRICE","FACTORY PACK PRICE / FACTORY YARD PRICE"])
            # Rhinestone Banding's price column is a combined text cell, e.g.
            # "$80.88 = $6.74/YARD" -- the actual variant price is just the first
            # figure; the second (per-yard) is used in description bullets, read
            # directly from the sheet when writing those, not stored as its own field.
            if clean(factory_price) and "=" in clean(factory_price):
                m = re.findall(r"[\d,]+\.?\d*", clean(factory_price))
                if m: factory_price = m[0].replace(",", "")
            factory_weight = get_first(row, ["WEIGHT PER FACTORY PACK","FACTORY PACK WEIGHT OZ"])
            mini_style = get_first(row, ["MINI PACK STYLE #","MINI PACK       STYLE #","MINI PACK STYLE NUMBER"])
            mini_qty = get_first(row, ["MINI PACK QUANTITY","MINI PACK QUANTITY DESCRIPTION","APPROXIMATE MINI PACK UNIT QUANTITY","MINI PACK UNIT QUANTITY"])
            mini_price = get_first(row, ["UNIT PRICE PER MINI PACK","MINI PACK PRICE"])
            mini_weight = get_first(row, ["WEIGHT PER MINI PACK","MINI PACK WEIGHT OZ"])
            # Neil's package-dimensions column, added 2026-08-04: a single free-text
            # cell like "10 x 5 x 3 in" rather than 3 separate Length/Width/Height
            # columns (unlike the earlier one-off Calcurates dimensions spreadsheet).
            # Kept as raw text, not split into numeric fields -- no assumptions about
            # which pack (Factory vs Mini) this describes until confirmed against a
            # real sheet; candidate list here is a starting guess, adjust once the
            # actual header text is seen.
            dimensions = get_first(row, ["DIMENSIONS","DIMENSIONS (L X W X H)","DIMENSIONS L X W X H","PACKAGE DIMENSIONS","BOX DIMENSIONS","PRODUCT DIMENSIONS"])
            inf = infer(ws.title, color_name, os.path.basename(path), forced_type_id=forced_type_id)
            if not inf["collection"]:
                raise ValueError(f"Could not detect a known product category for {os.path.basename(path)!r} (sheet {ws.title!r}, row {row_num}). Add a matching rule to infer() before importing this file.")
            title_descriptor = clean(shape) or inf["subcategory"]
            if inf.get("title_suffix"):
                # Collapse stray double-spacing typos from the source sheet, and
                # drop a redundant trailing "BEADS" from the color name itself
                # (e.g. "... LOOSE BEADS") when the suffix already ends in
                # "BEADS" -- avoids "LOOSE BEADS PRECIOSA ORNELA BEADS".
                color_display = re.sub(r"\s+", " ", clean(color_name)).strip()
                if inf["title_suffix"].split()[-1].upper() == "BEADS" and re.search(r"\bBEADS$", color_display, re.I):
                    color_display = re.sub(r"\s*\bBEADS$", "", color_display, flags=re.I).strip()
                title = " ".join(x for x in [inf["title_prefix"], clean(size), color_display, inf["title_suffix"]] if x)
            else:
                title = title_for(inf["title_prefix"], title_descriptor, size, color_name)
            bead_shape_value = clean(shape) if inf["populate_bead_shape_from_shape"] and clean(shape) else inf["bead_shape"]
            type_value = inf.get("type", "")
            row_text = " ".join(clean(x) for x in vals).lower()
            notes = []
            if "missing phot" in row_text or "mising phot" in row_text: notes.append("Source note: missing photo")
            if clean(factory_price) and not money(factory_price): notes.append(f"Invalid factory price in source: {clean(factory_price)!r}")
            if clean(mini_price) and not money(mini_price): notes.append(f"Invalid mini price in source: {clean(mini_price)!r}")
            # Compare only the size's leading token (e.g. "#2" out of "#2 (4.5mm)") against the
            # SKU -- some categories' SKUs only embed the short size class, not the full
            # parenthetical detail, so checking the whole string produced false positives.
            size_token = clean(size).split("(")[0].strip()
            if clean(mini_style) and size_token and size_token.lower() not in clean(mini_style).lower():
                notes.append(f"Mini pack style # may not match this row's size ({clean(size)!r}): {clean(mini_style)!r}")
            p = {
                "id":str(uuid.uuid4()), "source_file":os.path.basename(path), "source_sheet":ws.title, "source_row":row_num,
                "spreadsheet_type_id":inf.get("type_id",""),
                "approved":False, "skipped":False, "status":"Needs Review",
                "title":title, "handle":slugify(title), "brand":"", "vendor":inf.get("vendor") or "Margola",
                "collection":inf["collection"], "subcategory":inf["subcategory"], "color_type":clean(color_type_explicit) or inf["color_type"], "bead_shape":bead_shape_value, "type":type_value,
                "shopify_category":inf["shopify_category"],
                "size":clean(size), "size_mm":clean(size_mm), "color_number":clean(color_number), "color_name":clean(color_name),
                "image_filename":image_for(factory_style,color_name,title_descriptor,size), "image_src":"", "image_alt":title,
                "factory_style":clean(factory_style), "factory_quantity":clean(factory_qty),
                "factory_quantity_description":clean(factory_qty_desc) or clean(factory_qty) or inf["factory_qty_standard"],
                "factory_price":money(factory_price), "factory_weight_oz":weight_oz(factory_weight),
                "mini_style":clean(mini_style), "mini_quantity":clean(mini_qty) or inf["mini_qty_standard"],
                "mini_price":money(mini_price), "mini_weight_oz":weight_oz(mini_weight),
                "dimensions":clean(dimensions),

                "variants": build_variants(
                    factory_style=factory_style,
                    factory_price=factory_price,
                    factory_weight=factory_weight,
                    factory_qty_desc=clean(factory_qty_desc) or clean(factory_qty) or inf["factory_qty_standard"],
                    mini_style=mini_style,
                    mini_price=mini_price,
                    mini_weight=mini_weight,
                    mini_qty=clean(mini_qty) or inf["mini_qty_standard"],
                ),

                "source_description":clean(source_description), "generated_description":"", "description_approved":False,
                "notes":"; ".join(notes),
            }
            products.append(p)
    handles, skus = {}, {}
    for p in products:
        handles.setdefault(p["handle"], []).append(p["id"])
        for v in p.get("variants", []):
            sku = v.get("sku")
            if sku:
                skus.setdefault(sku, []).append(p["id"])
    dup_handles = {k for k,v in handles.items() if len(v)>1}
    dup_skus = {k for k,v in skus.items() if len(v)>1}
    for p in products:
        more = []
        if p["handle"] in dup_handles: more.append("Duplicate handle")
        for v in p.get("variants", []):
            if v.get("sku") in dup_skus:
                more.append(f"Duplicate variant SKU: {v.get('sku')}")
        if more: p["notes"] = "; ".join([x for x in [p["notes"]] + more if x])
        p["validation"] = validate(p)
        p["status"] = "Needs Review" if p["validation"]["warnings"] else "Ready for Review"
    return products

def load_products():
    if not os.path.exists(PRODUCTS_PATH): return []
    with open(PRODUCTS_PATH, "r", encoding="utf-8") as f: return json.load(f)

def save_products(products):
    for p in products: p["validation"] = validate(p)
    with open(PRODUCTS_PATH, "w", encoding="utf-8") as f: json.dump(products, f, indent=2)

def summary(products):
    return {"total":len(products),"approved":sum(p.get("approved") for p in products),"skipped":sum(p.get("skipped") for p in products),"needs_review":sum(1 for p in products if not p.get("approved") and not p.get("skipped")),"with_warnings":sum(1 for p in products if p.get("validation",{}).get("warnings")),"ready_100":sum(1 for p in products if p.get("validation",{}).get("score")==100)}

def detected_categories(products):
    counts = {}
    for p in products:
        key = (p.get("collection",""), p.get("subcategory",""))
        counts[key] = counts.get(key, 0) + 1
    return [{"collection":k[0],"subcategory":k[1],"count":v} for k,v in sorted(counts.items())]

def review_csv(products):
    fields = ["status","approved","skipped","validation_score","warnings","title","handle","collection","subcategory","color_type","bead_shape","type","size","color_number","color_name","image_filename","image_src","factory_style","factory_price","factory_weight_oz","mini_style","mini_price","mini_weight_oz","dimensions","source_sheet","source_row"]
    out = io.StringIO(); w = csv.DictWriter(out, fieldnames=fields); w.writeheader()
    for p in products:
        v = p.get("validation",{})
        row = {f:p.get(f,"") for f in fields}; row["validation_score"] = v.get("score",""); row["warnings"] = "; ".join(v.get("warnings",[]))
        w.writerow(row)
    return out.getvalue()


def oz_to_grams(value):
    try:
        return str(round(float(value) * 28.3495, 2))
    except:
        return ""

def shopify_rows(products, approved_only=True, limit=None, resolver=None):
    if resolver is None:
        resolver = MetaobjectResolver()
    rows, count = [], 0

    for p in products:
        if approved_only and (not p.get("approved") or p.get("skipped")):
            continue

        variants = p.get("variants", [])

        # Backward compatibility for old imported data without variants yet.
        if not variants:
            variants = build_variants(
                factory_style=p.get("factory_style"),
                factory_price=p.get("factory_price"),
                factory_weight=p.get("factory_weight_oz"),
                factory_qty_desc=p.get("factory_quantity_description"),
                mini_style=p.get("mini_style"),
                mini_price=p.get("mini_price"),
                mini_weight=p.get("mini_weight_oz"),
                mini_qty=p.get("mini_quantity"),
            )

        if not variants:
            continue

        if limit is not None and count >= limit:
            break

        count += 1
        body = p.get("generated_description") or p.get("source_description") or ""

        common = {
            "Handle": p["handle"],
            "Vendor": p.get("vendor") or "Margola",
            "Product Category": p.get("shopify_category", ""),
            "Type": p["subcategory"] or p.get("collection", ""),
            "Tags": ", ".join(x for x in [p["collection"], p["subcategory"], p["color_type"], p["bead_shape"], p["size"]] if x),
            "Published": "TRUE",
            "Variant Inventory Tracker": "shopify",
            "Variant Inventory Qty": 15,
            "Variant Inventory Policy": "deny",
            "Variant Fulfillment Service": "manual",
            "Variant Requires Shipping": "TRUE",
            "Variant Taxable": "TRUE",
            "Variant Weight Unit": "oz",
            "Status": "active"
        }

        for idx, variant in enumerate(variants):
            row = dict(common)
            row.update({
                "Title": p["title"] if idx == 0 else "",
                "Body (HTML)": body if idx == 0 else "",
                "Option1 Name": "Bundle Pack Options" if idx == 0 else "",
                "Option1 Value": variant.get("option_value", ""),
                "Option1 Linked To": "product.metafields.custom.bundle_pack_options" if idx == 0 else "",
                "Bundle Pack Options (product.metafields.custom.bundle_pack_options)": "; ".join(v["option_value"] for v in variants) if idx == 0 else "",
                "Color (product.metafields.shopify.color-pattern)": resolver.color_handle(p.get("color_name")) if idx == 0 else "",
                "Color Type (product.metafields.custom.color_type)": clean(p.get("color_type")) if idx == 0 else "",
                "Bead Size (product.metafields.custom.bead_size_mm)": (clean(p.get("size_mm")) or clean(p.get("size"))) if idx == 0 else "",
                "Size (product.metafields.shopify.size)": resolver.size_handle(p.get("size")) if idx == 0 else "",
                "Bead Shape (product.metafields.custom.bead_shape)": resolver.bead_shape_label(p.get("bead_shape")) if idx == 0 else "",
                "Type (product.metafields.custom.type)": resolver.type_label(p.get("type")) if idx == 0 else "",
                "Dimensions (product.metafields.custom.dimensions)": clean(p.get("dimensions","")) if idx == 0 else "",
                "Variant SKU": variant.get("sku", ""),
                "Variant Price": variant.get("price", ""),
                "Variant Grams": oz_to_grams(variant.get("weight_oz", "")),
                "Image Src": p.get("image_src", "") if idx == 0 else "",
                "Image Alt Text": (p.get("image_alt") or make_alt_text(p["title"])) if idx == 0 and p.get("image_src") else "",
            })
            rows.append(row)

    return rows

def shopify_csv(products, approved_only=True, limit=None):
    fields = [
        "Handle","Title","Body (HTML)","Vendor","Product Category","Type","Tags","Published",
        "Option1 Name","Option1 Value","Option1 Linked To",
        "Variant SKU","Variant Price","Variant Grams","Variant Weight Unit",
        "Variant Inventory Tracker","Variant Inventory Qty","Variant Inventory Policy",
        "Variant Fulfillment Service","Variant Requires Shipping","Variant Taxable",
        "Image Src","Image Alt Text",
        "Bundle Pack Options (product.metafields.custom.bundle_pack_options)",
        "Color (product.metafields.shopify.color-pattern)",
        "Color Type (product.metafields.custom.color_type)",
        "Bead Size (product.metafields.custom.bead_size_mm)",
        "Size (product.metafields.shopify.size)",
        "Bead Shape (product.metafields.custom.bead_shape)",
        "Type (product.metafields.custom.type)",
        "Status"
    ]

    out = io.StringIO()
    w = csv.DictWriter(out, fieldnames=fields, extrasaction="ignore")
    w.writeheader()

    resolver = MetaobjectResolver()
    for row in shopify_rows(products, approved_only, limit, resolver=resolver):
        w.writerow(row)

    return out.getvalue()

def parse_descriptions(text):
    blocks, cur, lines = {}, None, []
    for line in text.splitlines():
        m = re.match(r"^\s*===\s*(.*?)\s*===\s*$", line)
        if m:
            if cur: blocks[cur] = "\n".join(lines).strip()
            cur, lines = m.group(1).strip(), []
        else: lines.append(line)
    if cur: blocks[cur] = "\n".join(lines).strip()
    return {k:v for k,v in blocks.items() if k and v}

def guess(path):
    if path.endswith(".css"): return "text/css; charset=utf-8"
    if path.endswith(".js"): return "application/javascript; charset=utf-8"
    if path.endswith(".html"): return "text/html; charset=utf-8"
    return "application/octet-stream"


# --- Shopify client-credentials integration ---
def load_env():
    env_path = os.path.join(APP_DIR, ".env")
    if os.path.exists(env_path):
        for line in open(env_path, encoding="utf-8"):
            line=line.strip()
            if line and not line.startswith("#") and "=" in line:
                k,v=line.split("=",1)
                os.environ.setdefault(k.strip(), v.strip().strip('"').strip("'"))
load_env()
SHOPIFY_TOKEN_CACHE={"token":"","expires_at":0}
def shopify_config():
    return {"store":os.environ.get("SHOPIFY_STORE","").strip().replace("https://","").replace("http://","").strip("/"),"client_id":os.environ.get("SHOPIFY_CLIENT_ID","").strip(),"client_secret":os.environ.get("SHOPIFY_CLIENT_SECRET","").strip()}
def shopify_get_token(force=False):
    cfg=shopify_config(); now=time.time()
    if not cfg["store"] or not cfg["client_id"] or not cfg["client_secret"]: raise RuntimeError("Missing SHOPIFY_STORE, SHOPIFY_CLIENT_ID, or SHOPIFY_CLIENT_SECRET in .env")
    if not force and SHOPIFY_TOKEN_CACHE["token"] and SHOPIFY_TOKEN_CACHE["expires_at"]>now+60: return SHOPIFY_TOKEN_CACHE["token"]
    url=f"https://{cfg['store']}/admin/oauth/access_token"
    data=urllib.parse.urlencode({"grant_type":"client_credentials","client_id":cfg["client_id"],"client_secret":cfg["client_secret"]}).encode()
    req=urllib.request.Request(url,data=data,headers={"Content-Type":"application/x-www-form-urlencoded"},method="POST")
    with urllib.request.urlopen(req, timeout=30, context=ssl.create_default_context(cafile=certifi.where())) as resp: payload=json.loads(resp.read().decode())
    token=payload.get("access_token")
    if not token: raise RuntimeError(f"No access_token returned: {payload}")
    SHOPIFY_TOKEN_CACHE["token"]=token; SHOPIFY_TOKEN_CACHE["expires_at"]=now+int(payload.get("expires_in",86399)); return token
def shopify_graphql(query, variables=None):
    cfg=shopify_config(); token=shopify_get_token(); url=f"https://{cfg['store']}/admin/api/2026-04/graphql.json"
    body=json.dumps({"query":query,"variables":variables or {}}).encode()
    req=urllib.request.Request(url,data=body,headers={"Content-Type":"application/json","X-Shopify-Access-Token":token},method="POST")
    with urllib.request.urlopen(req, timeout=60, context=ssl.create_default_context(cafile=certifi.where())) as resp: payload=json.loads(resp.read().decode())
    if payload.get("errors"): raise RuntimeError(json.dumps(payload["errors"]))
    return payload["data"]
def shopify_test_connection(): return shopify_graphql("{ shop { name myshopifyDomain } }")["shop"]
def shopify_list_files():
    # Paginates through every file in the store (not just the first page) --
    # this store has thousands of files across other product lines, so a
    # single first-page fetch missed newly uploaded photos entirely.
    q="""query GetFiles($after:String){ files(first:250, after:$after, sortKey:CREATED_AT, reverse:true){ edges{ cursor node{ id alt createdAt ... on MediaImage{ image{ url } } ... on GenericFile{ url } } } pageInfo{ hasNextPage } } }"""
    out, cursor = [], None
    while True:
        data=shopify_graphql(q,{"after":cursor})
        edges=data["files"]["edges"]
        for e in edges:
            n=e["node"]; url=""
            if n.get("image"): url=n["image"].get("url") or ""
            if not url: url=n.get("url") or ""
            fn=urllib.parse.unquote(url.split("?")[0].rstrip("/").split("/")[-1]) if url else ""
            out.append({"id":n.get("id",""),"url":url,"filename":fn,"alt":n.get("alt",""),"createdAt":n.get("createdAt","")})
        if not data["files"]["pageInfo"]["hasNextPage"] or not edges: break
        cursor=edges[-1]["cursor"]
    return out
def file_map_path(): return os.path.join(DATA_DIR,"shopify_files.json")
def save_file_map(files): open(file_map_path(),"w",encoding="utf-8").write(json.dumps(files,indent=2))
def load_file_map():
    return json.load(open(file_map_path(),encoding="utf-8")) if os.path.exists(file_map_path()) else []
def apply_file_matches_to_products():
    products=load_products(); files=load_file_map(); by={f.get("filename","").lower():f.get("url") for f in files if f.get("filename") and f.get("url")}; matched=0
    valid_urls=set(by.values())
    for p in products:
        key=(p.get("image_filename") or "").lower()
        new_url=by.get(key)
        if not new_url and p.get("bead_shape")=="Roller Beads":
            # Roller Beads only shoot one photo per color and reuse it for both
            # the 6mm and 9mm listing, so fall back to the other size's filename.
            alt_key=""
            if "roller-6mm-" in key: alt_key=key.replace("roller-6mm-","roller-9mm-")
            elif "roller-9mm-" in key: alt_key=key.replace("roller-9mm-","roller-6mm-")
            if alt_key: new_url=by.get(alt_key)
        if new_url:
            p["image_src"]=new_url; matched+=1
        elif p.get("image_src") and p["image_src"] not in valid_urls:
            # A previously-matched file no longer exists on Shopify (e.g. it
            # was one of the test uploads deleted earlier) -- clear the
            # stale link instead of leaving a broken image_src that silently
            # excludes the product from every future drive scan.
            p["image_src"]=""
    save_products(products); return matched,len(products)

DRIVE_FILELISTS = [
    os.path.join(os.path.expanduser("~"), "Downloads", "harddrive_filelist.txt"),
    os.path.join(os.path.expanduser("~"), "Downloads", "2cut_filelist.txt"),
    os.path.join(os.path.expanduser("~"), "Downloads", "2cut_10_0_filelist.txt"),
    os.path.join(os.path.expanduser("~"), "Downloads", "bugle_filelist.txt"),
    os.path.join(os.path.expanduser("~"), "Downloads", "glass-jewels_filelist.txt"),
    os.path.join(os.path.expanduser("~"), "Downloads", "sew-on-glass-jewels_filelist.txt"),
    os.path.join(os.path.expanduser("~"), "Downloads", "cameos_filelist.txt"),
    os.path.join(os.path.expanduser("~"), "Downloads", "lochrosens_filelist.txt"),
    os.path.join(os.path.expanduser("~"), "Downloads", "rhinestone balls_filelist.txt"),
    os.path.join(os.path.expanduser("~"), "Downloads", "rhinestonebanding_filelist.txt"),
    os.path.join(os.path.expanduser("~"), "Downloads", "pearls-on-eye-pins_filelist.txt"),
    os.path.join(os.path.expanduser("~"), "Downloads", "leaf-bail_filelist.txt"),
]

# Hand-verified factory_style -> filename mapping for the one folder that
# holds (almost) every correctly-sized Roller Beads web photo, built by
# manually cross-referencing every file in the folder against every product's
# color number/name -- not the fuzzy RB-code guessing used elsewhere, since
# that guessing has repeatedly picked wrong-sized or wrong-variant (matte vs
# not) photos when this folder alone was already unambiguous. The color
# code collision between ROLLER-6MM-40010 (SMOKE GREY TRANSPARENT) and
# ROLLER-9MM-40010 (BLACK DIAMOND TRANSPARENT) turned out to be a
# spreadsheet error -- the 6mm row's code reverted back to its original
# 40020, leaving RB-4001.jpg unambiguously the 9mm product's photo, and
# RB-4002.jpg was added later as the 6mm product's own photo. Two products
# are still deliberately left out, waiting on new photos:
#  - ROLLER-9MM-1006M: no matte-specific photo exists in this folder at all.
#  - ROLLER-6MM-30040: no plain (non-matte) photo exists, only the 3004
#    matte one.
JPEG_FOR_WEB_FOLDER = "/Volumes/Hard Drive/Margola Import Corp./Product Images/Roller Beads/9MM/RB-WebImages-1-8-2020 3/JPEG for Web"
JPEG_FOR_WEB_MAP = {
    "ROLLER-6MM-00030": "RB-0002.jpg",
    "ROLLER-6MM-01790": "RB-red-metallic-01790.jpg",
    "ROLLER-6MM-03050": "RB-0300.jpg",
    "ROLLER-6MM-10020": "RB-1002.jpg",
    "ROLLER-6MM-10090": "10090.jpg",
    "ROLLER-6MM-10220": "10220.jpg",
    "ROLLER-6MM-10230": "RB-1023.jpg",
    "ROLLER-6MM-14400": "RB-14400.jpg",
    "ROLLER-6MM-20030": "RB-2003.jpg",
    "ROLLER-6MM-20050": "20050.jpg",
    "ROLLER-6MM-21415": "RB21415.jpg",
    "ROLLER-6MM-21435": "RB21435.jpg",
    "ROLLER-6MM-21455": "RB21455.jpg",
    "ROLLER-6MM-21495": "RB21495.jpg",
    "ROLLER-6MM-2398AB": "RB-2398.jpg",
    "ROLLER-6MM-30030": "RB-3003.jpg",
    "ROLLER-6MM-30060": "RB-3006.jpg",
    "ROLLER-6MM-30090": "30090.jpg",
    "ROLLER-6MM-31010": "31010.jpg",
    "ROLLER-6MM-40020": "RB-4002.jpg",
    "ROLLER-6MM-50130": "RB-5013.jpg",
    "ROLLER-6MM-50150": "RB-5015.jpg",
    "ROLLER-6MM-50250": "RB-5025.jpg",
    "ROLLER-6MM-5043M": "5043M.jpg",
    "ROLLER-6MM-50710": "50710.jpg",
    "ROLLER-6MM-50730": "RB-5073.jpg",
    "ROLLER-6MM-60010": "RB-6001.jpg",
    "ROLLER-6MM-60040": "RB-6004.jpg",
    "ROLLER-6MM-60080": "60080.jpg",
    "ROLLER-6MM-6302M": "6302M.jpg",
    "ROLLER-6MM-70100": "RB-7010.jpg",
    "ROLLER-6MM-80020": "RB-8002.jpg",
    "ROLLER-6MM-83110": "RB-8311---Copy.jpg",
    "ROLLER-6MM-90040": "RB-9004.jpg",
    "ROLLER-6MM-9008M": "9008M.jpg",
    "ROLLER-6MM-93140": "93140.jpg",
    "ROLLER-6MM-93210": "RB-9321.jpg",
    "ROLLER-6MM10060": "RB-1006.jpg",
    "ROLLER-9MM-00030": "RB-0002.jpg",
    "ROLLER-9MM-01700": "RB-01700-silver-metallic.jpg",
    "ROLLER-9MM-01710": "RB-gold-metallic-01710.jpg",
    "ROLLER-9MM-01720": "RB-DK-brown-metallic-01720.jpg",
    "ROLLER-9MM-01730": "RB-blue-metallic-01730.jpg",
    "ROLLER-9MM-01750": "RB-DK-green-metallic-01750.jpg",
    "ROLLER-9MM-01790": "RB-red-metallic-01790.jpg",
    "ROLLER-9MM-02010": "RB-02010.jpg",
    "ROLLER-9MM-03050": "RB-0300.jpg",
    "ROLLER-9MM-10020": "RB-1002.jpg",
    "ROLLER-9MM-10060": "RB-1006.jpg",
    "ROLLER-9MM-10090": "10090.jpg",
    "ROLLER-9MM-1011M": "RB-1011m.jpg",
    "ROLLER-9MM-10220": "10220.jpg",
    "ROLLER-9MM-10230": "RB-1023.jpg",
    "ROLLER-9MM-10250": "RB-1025.jpg",
    "ROLLER-9MM-13600": "RB-13600.jpg",
    "ROLLER-9MM-13620": "RB-13620.jpg",
    "ROLLER-9MM-14400": "RB-14400.jpg",
    "ROLLER-9MM-18016": "RB-18016.jpg",
    "ROLLER-9MM-20030": "RB-2003.jpg",
    "ROLLER-9MM-2003M": "RB-2003m.jpg",
    "ROLLER-9MM-20060": "roller-9mm-20060.jpg",
    "ROLLER-9MM-20080": "RB-2008.jpg",
    "ROLLER-9MM-21415": "RB21415.jpg",
    "ROLLER-9MM-21435": "RB21435.jpg",
    "ROLLER-9MM-23030": "RB-2303.jpg",
    "ROLLER-9MM-23980": "RB-2398.jpg",
    "ROLLER-9MM-30030": "RB-3003.jpg",
    "ROLLER-9MM-3004M": "RB-30040-matte.jpg",
    "ROLLER-9MM-30060": "RB-3006.jpg",
    "ROLLER-9MM-30090": "30090.jpg",
    "ROLLER-9MM-31010": "31010.jpg",
    "ROLLER-9MM-32010": "32010.jpg",
    "ROLLER-9MM-33060": "RB-3306.jpg",
    "ROLLER-9MM-40010": "RB-4001.jpg",
    "ROLLER-9MM-4302L": "RB-14069.jpg",
    "ROLLER-9MM-50130": "RB-5013.jpg",
    "ROLLER-9MM-50150": "RB-5015.jpg",
    "ROLLER-9MM-50230": "RB-5023.jpg",
    "ROLLER-9MM-50250": "RB-5025.jpg",
    "ROLLER-9MM-50710": "50710.jpg",
    "ROLLER-9MM-50730": "RB-5073.jpg",
    "ROLLER-9MM-5073M": "RB-5073-matte.jpg",
    "ROLLER-9MM-51010": "51010.jpg",
    "ROLLER-9MM-52120": "52120.jpg",
    "ROLLER-9MM-60010": "RB-6001.jpg",
    "ROLLER-9MM-60040": "RB-6004.jpg",
    "ROLLER-9MM-60080": "60080.jpg",
    "ROLLER-9MM-60150": "RB-6015.jpg",
    "ROLLER-9MM-63040": "63040.jpg",
    "ROLLER-9MM-63110": "RB-6311.jpg",
    "ROLLER-9MM-63130": "RB-6313.jpg",
    "ROLLER-9MM-70100": "RB-7010.jpg",
    "ROLLER-9MM-70110": "RB-7011.jpg",
    "ROLLER-9MM-81210": "81210.jpg",
    "ROLLER-9MM-83110": "RB-8311---Copy.jpg",
    "ROLLER-9MM-90040": "RB-9004.jpg",
    "ROLLER-9MM-90050": "RB-9005.jpg",
    "ROLLER-9MM-90070": "RB-9007.jpg",
    "ROLLER-9MM-90090": "RB-9009.jpg",
    "ROLLER-9MM-90120": "RB-9012.jpg",
    "ROLLER-9MM-93110": "93110.jpg",
    "ROLLER-9MM-93140": "93140.jpg",
    "ROLLER-9MM-93200": "RB-9320---Copy.jpg",
    "ROLLER-9MM-93210": "RB-9321.jpg",
    "ROLLER-9MM71010": "RB-7101.jpg",
    "ROLLER-9MM73010": "73010.jpg",
}

# Same approach as JPEG_FOR_WEB_MAP: hand-verified factory_style -> filename,
# built by cross-referencing PLEXY-LALIQUE FLOWERS AND LEAVES (1).xlsx against
# every file in the folder. PL1634-27mm-Petals has no photo anywhere on the
# drive at all and is deliberately left out. PL1646 was 17mm in the
# spreadsheet but the only photo on the drive was named 27mm -- Yair
# confirmed the photo was renamed to 17mm to match, so this assumes the
# renamed file kept the same "PL1646-17mm-DaisyButtons.jpg" pattern as its
# neighbors (verify this filename is right after pulling the actual drive
# listing, since it wasn't independently confirmed).
PLEXI_LALIQUE_FOLDER = "/Volumes/Hard Drive/Margola Import Corp./Product Images/Plexi Lalique Flowers &Leaves/JPEGS"
PLEXI_LALIQUE_MAP = {
    "PL1631-20mm-Flower": "PL1631-20mm-Flower.jpg",
    "PL1632-24mm-Petals": "PL1632-24mm-Petals.jpg",
    "PL1637-23mm-Daisy": "PL1637-23mm-Daisy.jpg",
    "PL1638-26mm-Petals": "PL1638-26mm-Petals.jpg",
    "PL1640-34mm-Petals": "PL1640-34x33mm-Petals.jpg",
    "PL1645-23mm-Daisy button": "PL1645-23mm-DaisyButtons.jpg",
    "PL1646-17mm-Daisy button": "PL1646-17mm-DaisyButtons.jpg",
    "PL1647-17mm-Petals": "PL1647-17mm-Petals.jpg",
    "PL1648-25x15mm-Leaf": "PL1648-25x15mm-Leaf.jpg",
    "PL1649-18x24mm-Leaf": "PL1649-18x24mm-Leaf.jpg",
    "PL1650-27x26mm-Leaf": "PL1650-27x26mm-Leaf.jpg",
    "PL1651-21x20mm-Leaf": "PL1651-21x20mm-Leaf.jpg",
    "PL1653-34x32mm-Leaves": "PL1653-34x32mm-Leaves.jpg",
    "PL1654-52.5x15mm-Leaf": "PL1654--52.5x15mm-Leaf.jpg",
    "PL1655-41x15mm-Leaf": "PL1655-41x15mm-Leaf.jpg",
    "PL1656-30x17mm-Petals": "PL1656-30x17mm-Petals.jpg",
    "PL1657-27X10mm-Leaf": "PL1657-27x10mm-Leaf.jpg",
}

# Registry of every hand-verified "one folder, one map" photo source, so the
# scan endpoint/button can be reused for each new category instead of
# duplicating an endpoint per folder.
PHOTO_MAPS = {
    "roller_crow_jpeg_for_web": {"label": "Roller/Crow JPEGs for Web", "folder": JPEG_FOR_WEB_FOLDER, "map": JPEG_FOR_WEB_MAP},
    "plexi_lalique": {"label": "Plexi-Lalique Flowers & Leaves", "folder": PLEXI_LALIQUE_FOLDER, "map": PLEXI_LALIQUE_MAP},
}

JUNK_PATH_MARKERS = ("._", ".DS_Store", ".psd")
# Deprioritized rather than excluded outright -- a couple of the drive's
# only correctly-sized "JPEG for Web" copies are themselves named with
# "---Copy" (e.g. RB-8311---Copy.jpg), so treating it as junk discarded
# the one usable file instead of an actual redundant duplicate.
COPY_MARKERS = (" - copy", "---copy")

GOOD_SUBFOLDER_MARKERS = ("jpeg for web", "jpegs for web", "jpegs", "jpeg")
TARGET_IMAGE_SIZE = (1000, 1000)

def _image_dimensions(path):
    # Reads just enough of the file header to get pixel dimensions -- no
    # Pillow/other dependency needed, and cheap even for a large batch.
    try:
        with open(path, "rb") as f:
            head = f.read(2)
            if head == b"\xff\xd8":  # JPEG
                while True:
                    marker = f.read(2)
                    if len(marker) < 2 or marker[0] != 0xFF: return None
                    if marker[1] in (0xC0,0xC1,0xC2,0xC3,0xC5,0xC6,0xC7,0xC9,0xCA,0xCB,0xCD,0xCE,0xCF):
                        f.read(3)
                        h = int.from_bytes(f.read(2), "big")
                        w = int.from_bytes(f.read(2), "big")
                        return (w, h)
                    seg_len = int.from_bytes(f.read(2), "big")
                    if seg_len < 2: return None
                    f.seek(seg_len - 2, 1)
            elif head == b"\x89P":  # PNG
                f.seek(16)
                w = int.from_bytes(f.read(4), "big")
                h = int.from_bytes(f.read(4), "big")
                return (w, h)
    except Exception:
        return None
    return None

def _pick_best_candidate(paths):
    # When the same filename exists in more than one place (e.g. a parent
    # folder with full-res 2000x2000 originals and a "JPEG for Web"
    # subfolder with the correct 1000x1000 versions), prefer the one that's
    # actually the right size, falling back to the "jpeg(s) for web"
    # subfolder naming convention if the drive isn't mounted to check.
    if len(paths) == 1: return paths[0]
    def score(path):
        dims = _image_dimensions(path) if os.path.exists(path) else None
        size_ok = dims == TARGET_IMAGE_SIZE
        folder_ok = any(m in path.lower() for m in GOOD_SUBFOLDER_MARKERS)
        not_a_copy = not any(m in path.lower() for m in COPY_MARKERS)
        return (size_ok, folder_ok, not_a_copy)
    return max(paths, key=score)

def load_drive_index():
    # Reads the pre-generated drive filelist dumps (tab-separated: size, mtime, path)
    # rather than scanning the drives live, since they aren't always plugged in.
    roller_9mm, roller_6mm, crow, leather_cord, two_cut, bugle, generic = {}, {}, {}, {}, {}, {}, {}
    found_any = False
    for list_path in DRIVE_FILELISTS:
        if not os.path.exists(list_path): continue
        found_any = True
        for line in open(list_path, encoding="utf-8", errors="replace").read().splitlines():
            parts = line.split("\t")
            if len(parts) < 3: continue
            full = parts[2]
            if any(m in full for m in JUNK_PATH_MARKERS): continue
            base = full.rsplit("/", 1)[-1].lower()
            fl = full.lower()
            generic.setdefault(base, []).append(full)
            if "/roller beads/9mm/" in fl: roller_9mm.setdefault(base, []).append(full)
            elif "/roller beads/6mm/" in fl: roller_6mm.setdefault(base, []).append(full)
            elif "/crow" in fl: crow.setdefault(base, []).append(full)
            elif "/leather cord/" in fl: leather_cord.setdefault(base, []).append(full)
            elif "/2 cuts/" in fl: two_cut.setdefault(base, []).append(full)
            elif "/bugle beads/" in fl: bugle.setdefault(base, []).append(full)
    roller_9mm_raw, roller_6mm_raw = roller_9mm, roller_6mm
    roller_9mm, roller_6mm, crow, leather_cord, two_cut, bugle, generic = (
        {base: _pick_best_candidate(paths) for base, paths in pool.items()}
        for pool in (roller_9mm, roller_6mm, crow, leather_cord, two_cut, bugle, generic)
    )
    # Built from the raw (pre-picked) path lists, not the already-resolved
    # per-size dicts above -- otherwise a basename shared by both sizes
    # (e.g. a 6MM full-res original and a 9MM "JPEG for Web" copy that
    # happen to share the same filename) collapses to whichever size was
    # merged in first, silently discarding a correctly-sized duplicate.
    roller_union = {
        base: _pick_best_candidate((roller_6mm_raw.get(base, []) + roller_9mm_raw.get(base, [])))
        for base in set(roller_6mm_raw) | set(roller_9mm_raw)
    }
    return {"found_any": found_any, "filelists": DRIVE_FILELISTS,
            "roller_9mm": roller_9mm, "roller_6mm": roller_6mm, "roller_union": roller_union,
            "crow": crow, "leather_cord": leather_cord, "two_cut": two_cut, "bugle": bugle, "generic": generic}

def _roller_candidates(color_number):
    cn = clean(color_number)
    cands = []
    m = re.search(r"use\s+(?:color\s+)?(?:the\s+)?([a-z0-9 .]+?)(?:\s+image\)?|\)|$)", cn, re.I)
    if m:
        hint = m.group(1).strip().lower().replace("metallc", "metallic")
        cands.append(hint.replace(" ", "-"))
        cands.append(re.sub(r"[^0-9a-z]", "", hint))
    raw = cn.split("(")[0].strip()
    is_matte = "matte" in raw.lower()
    digits = re.sub(r"[^0-9]", "", raw)
    if digits:
        variants = [digits]
        if len(digits) == 5 and digits.endswith("0"): variants.append(digits[:-1])
        for v in list(variants):
            # For a matte color, try the matte-suffixed filename before the
            # plain one -- otherwise, when a folder has both a matte and a
            # non-matte photo for the same base code, _find_all_roller_matches
            # adds the non-matte exact match first and _pick_best_candidate's
            # tie-break (same size/folder) silently keeps the wrong photo.
            if is_matte: cands.append(v + "m")
            cands.append(v)
    return [c for c in cands if c]

def _find_prefixed(pool, prefixes):
    for prefix in prefixes:
        for k, v in pool.items():
            if k.startswith(prefix): return v
    return None

def _find_all_roller_matches(pool, cands):
    # Collects every match across every candidate rather than stopping at
    # the first hit -- a spreadsheet's "(USE X)" cross-reference note and
    # its own raw color code can each independently match a *different*
    # real file (one a full-res original, one the correctly-sized web
    # copy), and whichever candidate string happens to be tried first
    # shouldn't decide the winner. The final size/folder-based scoring in
    # _pick_best_candidate chooses among all of them.
    results = []
    def add(v):
        if v and v not in results: results.append(v)
    for c in cands:
        for key in (f"rb-{c}.jpg", f"rb{c}.jpg", f"rb-{c}.jpeg", f"rb{c}.jpeg", f"{c}.jpg", f"{c}.jpeg"):
            add(pool.get(key))
    for c in cands:
        for prefix in (f"rb-{c}-", f"rb-{c} ", f"rb{c}-", f"rb-{c}.", f"rb{c}."):
            for k, v in pool.items():
                if k.startswith(prefix): add(v)
    for c in cands:
        if len(c) >= 4 and c.isalnum():
            for suffix in (f"-{c}.jpg", f"-{c}.jpeg"):
                for k, v in pool.items():
                    if k.endswith(suffix): add(v)
    return results

def _find_2cut_photo(pool, color_number):
    # Filenames are the color number plus the size marker, e.g. "35061-11_0.jpg"
    # (dash or underscore, both seen in practice). One product (48102) is
    # labeled "10_0" instead of "11_0" on the actual file -- matches the
    # Image File Name typo already present in Neil's own spreadsheet for that
    # row, so it's trusted as the real filename rather than treated as junk.
    # Pool keys are lowercased (load_drive_index lowercases every basename),
    # so the code must be too -- matte-suffix codes like "1107M" are the
    # first color numbers with letters in them, exposing this.
    code = clean(color_number).lower()
    for size in ("11_0", "10_0", "11", "10"):
        for sep in ("_", "-"):
            for ext in (".jpg", ".jpeg", ".png"):
                v = pool.get(f"{code}{sep}{size}{ext}")
                if v: return v
    return None

def _find_bugle_photo(pool, color_number, size_class):
    # Filenames are the color number plus the bugle size CLASS digit (e.g.
    # "05051_2.jpg" for the "#2" bugle line), not the mm diameter -- confirmed
    # against the real bugle_filelist.txt dump 2026-07-25 (every color in the
    # #2 batch has exactly one "{code}_2.jpg", plus macOS "._" shadow copies
    # already filtered out by JUNK_PATH_MARKERS).
    code = clean(color_number).lower()
    # size_class is the raw "size" field, e.g. "#2 (4.5mm)" -- only the leading
    # class number (before the parenthetical mm detail) appears in filenames.
    leading = clean(size_class).split()[0] if clean(size_class) else ""
    size = re.sub(r"[^0-9]", "", leading)
    if not size: return None
    for sep in ("_", "-"):
        for ext in (".jpg", ".jpeg", ".png"):
            v = pool.get(f"{code}{sep}{size}{ext}")
            if v: return v
    return None

def _leather_cord_candidate(image_filename):
    m = re.match(r"lc-(\d+)-(\d+)mm-(.+)\.jpg$", image_filename or "")
    if not m: return None
    whole, frac, color = m.groups()
    size = whole if frac == "0" else f"{whole}.{frac}"
    return size, color.replace("-", " ")

def _find_sew_on_glass_jewel_photo(pool, factory_style):
    # Generic exact-match already covers most of this category. Two known
    # gaps in the real photo folder: (1) a couple of photos were shot before
    # their SKU's "V" (Vintage) prefix got corrected in the sheet, so the
    # filename is missing it; (2) one filename has an extra alternate German
    # style number suffix baked in (e.g. "...--(04327-german-style-no.).jpg").
    code = clean(factory_style).lower()
    no_v = re.sub(r"^(\d+)v(-)", r"\1\2", code)
    v = pool.get(f"{no_v}.jpg")
    if v: return v
    for k, val in pool.items():
        if k.startswith(code + "--") or k.startswith(code + "-("):
            return val
    return None

def _find_cameos_photo(pool, factory_style):
    # Most of this category's real photo filenames were named directly from
    # the raw SKU text (spaces kept as spaces within multi-word color names,
    # not converted to dashes the way slugify() does), so try that literal
    # form before falling back further. A handful of files on disk also
    # predate later spreadsheet fixes: WHTE->WHITE / BOUIQUET->BOUQUET typos,
    # and one missing "V" (confirmed 2026-07-28 against cameos_filelist.txt).
    raw = clean(factory_style).lower()
    candidates = [raw]
    candidates.append(re.sub(r"^(\d+)v(-)", r"\1\2", raw))
    old_spelling = raw.replace("white", "whte").replace("bouquet", "bouiquet")
    candidates.append(old_spelling)
    candidates.append(re.sub(r"^(\d+)v(-)", r"\1\2", old_spelling))
    for c in candidates:
        v = pool.get(f"{c}.jpg") or pool.get(f"{c} .jpg")
        if v: return v
    return None

def _find_rhinestone_balls_photo(pool, color_number):
    # Real photos exist only per plating (Crystal/Gold or Crystal/Silver), not per
    # size -- confirmed 2026-07-31 against the real filelist: only two usable files,
    # "9000-0002G.jpg" and "9000-0002S.jpg", meant to be reused across all three
    # sizes (6mm/8mm/10mm) of that plating. Two other files in the same folder
    # (Crystal-AB-6MM.jpg, Crystal-Gold-6MM.jpg) don't match any product in this
    # sheet (no AB variant exists here) and were explicitly not the ones to use.
    code = clean(color_number).lower()
    return pool.get(f"9000-{code}.jpg")

def _find_metal_set_banding_photo(pool, factory_style):
    # Generic exact-match already covers every row except one: SKU 31501B has no
    # photo of its own on the drive -- Neil confirmed (2026-07-31) it should reuse
    # 31901B's photo (same one-row/black-net construction, different stone size).
    code = clean(factory_style).lower()
    if code.startswith("31501b"):
        return pool.get("31901b-19ss-0002s.jpg")
    return None

def _find_pearls_photo(pool, color_number, size):
    # Real photos are named "{color number}-{SIZE}.jpg", SIZE in uppercase mm
    # with an underscore standing in for the "x" in compound dimensions (e.g.
    # "70402-10_6MM.jpg" for the 10x6mm pear). One known drive-naming quirk:
    # the 18x6mm row's only real photo is filed as "16_8MM" -- Yair confirmed
    # 2026-08-04 the sheet's 18x6mm is the correct product size, so this is a
    # one-off filename reuse, not a size to fix in the sheet.
    code = clean(color_number).lower()
    s = clean(size).lower()
    if s == "18x6mm":
        v = pool.get(f"{code}-16_8mm.jpg")
        if v: return v
    size_key = s.replace("x", "_")
    for ext in (".jpg", ".jpeg", ".png"):
        v = pool.get(f"{code}-{size_key}{ext}")
        if v: return v
    return None

def _find_leaf_bail_photo(pool, factory_style):
    # Only two real photos exist, named by finish rather than SKU/color-code:
    # "GoldLeafBail.jpg" (GL) and "SilverLeafBail.jpg" (IR = Imitation Rhodium,
    # a silver-tone finish) -- confirmed against leaf-bail_filelist.txt 2026-08-04.
    code = clean(factory_style).lower()
    if code.endswith("-gl"): return pool.get("goldleafbail.jpg")
    if code.endswith("-ir"): return pool.get("silverleafbail.jpg")
    return None

def resolve_photo_from_drive(product, index):
    # Read-only lookup: finds a real file on the indexed drives for a product
    # missing image_src. Does not touch Shopify or the filesystem.
    target = clean(product.get("image_filename"))
    if not target: return None
    key = target.lower()
    # An exact match on the product's own expected filename (e.g. someone
    # already renamed/exported "roller-9mm-20060.jpg" to match our output
    # naming directly) is always correct -- check this before any of the
    # fuzzier per-category "RB-code" guessing below, since that guessing
    # can otherwise land on a wrong-sized file first and never get here.
    if key in index["generic"]:
        return {"source_path": index["generic"][key], "target_filename": target, "reused_other_size": False}
    if product.get("bead_shape") == "Crow Beads":
        code = re.sub(r"[^0-9]", "", product.get("color_number") or "")
        found = _find_prefixed(index["crow"], (f"crow-9mm-{code}-", f"crow-9mm-{code} "))
        if found: return {"source_path": found, "target_filename": target, "reused_other_size": False}
    elif product.get("bead_shape") == "Roller Beads":
        is_9mm = "9mm" in key
        own_pool = index["roller_9mm"] if is_9mm else index["roller_6mm"]
        cands = _roller_candidates(product.get("color_number") or "")
        own_matches = _find_all_roller_matches(own_pool, cands)
        union_matches = _find_all_roller_matches(index["roller_union"], cands)
        all_matches = own_matches + [m for m in union_matches if m not in own_matches]
        if all_matches:
            best = _pick_best_candidate(all_matches)
            return {"source_path": best, "target_filename": target, "reused_other_size": best not in own_matches}
    elif product.get("subcategory") == "Leather Cord":
        parsed = _leather_cord_candidate(target)
        if parsed:
            size, color = parsed
            for k, v in index["leather_cord"].items():
                kk = k.replace("-", " ")
                if f"{size}mm" in kk and color in kk:
                    return {"source_path": v, "target_filename": target, "reused_other_size": False}
    elif product.get("bead_shape") == "2 Cut Beads":
        found = _find_2cut_photo(index["two_cut"], product.get("color_number"))
        if found: return {"source_path": found, "target_filename": target, "reused_other_size": False}
    elif product.get("bead_shape") == "Bugle Beads":
        found = _find_bugle_photo(index["bugle"], product.get("color_number"), product.get("size"))
        if found: return {"source_path": found, "target_filename": target, "reused_other_size": False}
    elif product.get("spreadsheet_type_id") == "sew-on-glass-jewels":
        found = _find_sew_on_glass_jewel_photo(index["generic"], product.get("factory_style"))
        if found: return {"source_path": found, "target_filename": target, "reused_other_size": False}
    elif product.get("spreadsheet_type_id") == "cameos-intaglios":
        found = _find_cameos_photo(index["generic"], product.get("factory_style"))
        if found: return {"source_path": found, "target_filename": target, "reused_other_size": False}
    elif product.get("spreadsheet_type_id") == "rhinestone-balls":
        found = _find_rhinestone_balls_photo(index["generic"], product.get("color_number"))
        if found: return {"source_path": found, "target_filename": target, "reused_other_size": True}
    elif product.get("spreadsheet_type_id") == "metal-set-rhinestone-banding":
        found = _find_metal_set_banding_photo(index["generic"], product.get("factory_style"))
        if found: return {"source_path": found, "target_filename": target, "reused_other_size": True}
    elif product.get("spreadsheet_type_id") == "pearls-on-eye-pins":
        found = _find_pearls_photo(index["generic"], product.get("color_number"), product.get("size"))
        if found: return {"source_path": found, "target_filename": target, "reused_other_size": clean(product.get("size")).lower() == "18x6mm"}
    elif product.get("spreadsheet_type_id") == "leaf-bail":
        found = _find_leaf_bail_photo(index["generic"], product.get("factory_style"))
        if found: return {"source_path": found, "target_filename": target, "reused_other_size": False}
    return None

def multipart_form_data(fields, files):
    boundary="----MargolaBoundary"+uuid.uuid4().hex; chunks=[]
    for k,v in fields.items(): chunks += [f"--{boundary}\r\n".encode(), f'Content-Disposition: form-data; name="{k}"\r\n\r\n'.encode(), str(v).encode(), b"\r\n"]
    for k,(filename,ctype,data) in files.items(): chunks += [f"--{boundary}\r\n".encode(), f'Content-Disposition: form-data; name="{k}"; filename="{filename}"\r\n'.encode(), f"Content-Type: {ctype}\r\n\r\n".encode(), data, b"\r\n"]
    chunks.append(f"--{boundary}--\r\n".encode()); return boundary,b"".join(chunks)
def shopify_upload_file(filepath, alt="", filename=None):
    filename=filename or os.path.basename(filepath); ctype=mimetypes.guess_type(filepath)[0] or "image/jpeg"; size=str(os.path.getsize(filepath))
    q1="""mutation stagedUploadsCreate($input:[StagedUploadInput!]!){ stagedUploadsCreate(input:$input){ stagedTargets{ url resourceUrl parameters{ name value } } userErrors{ field message } } }"""
    d=shopify_graphql(q1,{"input":[{"filename":filename,"mimeType":ctype,"httpMethod":"POST","resource":"FILE","fileSize":size}]})["stagedUploadsCreate"]
    if d.get("userErrors"): raise RuntimeError(json.dumps(d["userErrors"]))
    t=d["stagedTargets"][0]; fields={p["name"]:p["value"] for p in t["parameters"]}; data=open(filepath,"rb").read(); boundary,body=multipart_form_data(fields,{"file":(filename,ctype,data)})
    req=urllib.request.Request(t["url"],data=body,headers={"Content-Type":f"multipart/form-data; boundary={boundary}"},method="POST")
    urllib.request.urlopen(req, timeout=120, context=ssl.create_default_context(cafile=certifi.where())).read()
    q2="""mutation fileCreate($files:[FileCreateInput!]!){ fileCreate(files:$files){ files{ id alt fileStatus createdAt ... on MediaImage{ image{ url } } ... on GenericFile{ url } } userErrors{ field message } } }"""
    c=shopify_graphql(q2,{"files":[{"alt":alt or filename,"contentType":"IMAGE","originalSource":t["resourceUrl"]}]})["fileCreate"]
    if c.get("userErrors"): raise RuntimeError(json.dumps(c["userErrors"]))
    f=c["files"][0]; url=(f.get("image") or {}).get("url") or f.get("url") or ""
    return {"id":f.get("id",""),"filename":filename,"url":url,"alt":f.get("alt",""),"status":f.get("fileStatus",""),"createdAt":f.get("createdAt","")}

def shopify_find_or_create_collection(title, handle):
    q="""query($handle:String!){ collectionByIdentifier(identifier:{handle:$handle}){ id title } }"""
    existing=shopify_graphql(q,{"handle":handle}).get("collectionByIdentifier")
    if existing: return existing["id"], False
    m="""mutation($input:CollectionInput!){ collectionCreate(input:$input){ collection{ id } userErrors{ field message } } }"""
    d=shopify_graphql(m,{"input":{"title":title,"handle":handle}})["collectionCreate"]
    if d.get("userErrors"): raise RuntimeError(json.dumps(d["userErrors"]))
    return d["collection"]["id"], True

def shopify_product_id_by_handle(handle):
    q="""query($handle:String!){ productByIdentifier(identifier:{handle:$handle}){ id } }"""
    p=shopify_graphql(q,{"handle":handle}).get("productByIdentifier")
    return p["id"] if p else None

def shopify_add_product_to_collection(product_id, collection_id):
    m="""mutation($product:ProductUpdateInput!){ productUpdate(product:$product){ product{ id } userErrors{ field message } } }"""
    d=shopify_graphql(m,{"product":{"id":product_id,"collectionsToJoin":[collection_id]}})["productUpdate"]
    if d.get("userErrors"): raise RuntimeError(json.dumps(d["userErrors"]))

# Only needed for collections that span more than one size (e.g. 2 Cut Beads'
# 10/0 + 11/0). Explicit map, not inferred -- extend per category as needed.
COLLECTION_SIZE_ORDER = {
    "2-cut-beads": {"10/0": 0, "11/0": 1},
}

def sort_key(sku, size_order=None, color_name_fallback="", group_by_size=True):
    parts = sku.split("-")
    size = parts[1] if len(parts) > 1 else ""
    color_raw = parts[2] if len(parts) > 2 else ""
    m = re.match(r"(\d+)", color_raw)
    if m:
        color_num = int(m.group(1))
        color_sort = (0, color_num, color_raw)          # numeric codes sort first, by number
    else:
        color_sort = (1, 0, color_name_fallback.upper()) # no numeric code -> alphabetical fallback
    if not group_by_size:
        return color_sort
    return ((size_order or {}).get(size, 99), *color_sort)

def shopify_collection_by_handle(handle):
    q="""query($handle:String!){ collectionByIdentifier(identifier:{handle:$handle}){ id title } }"""
    return shopify_graphql(q,{"handle":handle}).get("collectionByIdentifier")

def shopify_collection_products(collection_id):
    # Fetched with sortKey:COLLECTION_DEFAULT so the first page reflects the
    # *current* live manual order, for before/after comparison in the review UI.
    nodes, cursor = [], None
    while True:
        q="""query($id:ID!,$after:String){ collection(id:$id){ products(first:100, after:$after, sortKey:COLLECTION_DEFAULT){ edges{ cursor node{ id title handle variants(first:1){ edges{ node{ sku } } } } } pageInfo{ hasNextPage } } } }"""
        data=shopify_graphql(q,{"id":collection_id,"after":cursor})["collection"]["products"]
        edges=data["edges"]
        nodes.extend(e["node"] for e in edges)
        if not data["pageInfo"]["hasNextPage"] or not edges: break
        cursor=edges[-1]["cursor"]
    return nodes

def shopify_collection_sort_proposal(handle, group_by_size=True):
    collection = shopify_collection_by_handle(handle)
    if not collection:
        raise RuntimeError(f"Collection not found for handle {handle!r}")
    products = shopify_collection_products(collection["id"])
    size_order = COLLECTION_SIZE_ORDER.get(handle, {})
    current_order, enriched, sizes_present = [], [], set()
    for p in products:
        variant_edges = (p.get("variants") or {}).get("edges", [])
        sku = variant_edges[0]["node"]["sku"] if variant_edges else ""
        parts = sku.split("-")
        size = parts[1] if len(parts) > 1 else ""
        if size: sizes_present.add(size)
        current_order.append(p["title"])
        key = sort_key(sku, size_order, color_name_fallback=p.get("title",""), group_by_size=group_by_size)
        enriched.append({"id":p["id"], "title":p["title"], "sku":sku, "sort_key":key})
    enriched.sort(key=lambda e: e["sort_key"])
    return {
        "collection_id": collection["id"], "collection_title": collection["title"], "handle": handle,
        "multi_size": len(sizes_present) > 1, "sizes_present": sorted(sizes_present),
        "group_by_size": group_by_size,
        "current_order_titles": current_order,
        "products": [{"id":e["id"], "title":e["title"], "sku":e["sku"]} for e in enriched],
    }

def shopify_apply_collection_sort(collection_id, ordered_product_ids):
    set_manual_m = """mutation SetManual($input: CollectionInput!) { collectionUpdate(input: $input) { userErrors { field message } } }"""
    d1 = shopify_graphql(set_manual_m, {"input": {"id": collection_id, "sortOrder": "MANUAL"}})["collectionUpdate"]
    if d1.get("userErrors"): raise RuntimeError(json.dumps(d1["userErrors"]))

    reorder_m = """mutation Reorder($id: ID!, $moves: [MoveInput!]!) { collectionReorderProducts(id: $id, moves: $moves) { job { id done } userErrors { field message } } }"""
    moves = [{"id": pid, "newPosition": i} for i, pid in enumerate(ordered_product_ids)]
    d2 = shopify_graphql(reorder_m, {"id": collection_id, "moves": moves})["collectionReorderProducts"]
    if d2.get("userErrors"): raise RuntimeError(json.dumps(d2["userErrors"]))

    job = d2.get("job") or {}
    job_id, done = job.get("id"), job.get("done", False)
    if job_id and not done:
        for _ in range(30):
            time.sleep(1)
            jd = shopify_graphql('{ job(id: "%s") { done } }' % job_id).get("job") or {}
            if jd.get("done"):
                done = True; break

    # Step 3: verify against live data -- don't just trust a clean userErrors.
    verify_q = """query($id:ID!){ collection(id:$id){ sortOrder products(first:10, sortKey:COLLECTION_DEFAULT){ edges{ node{ id title } } } } }"""
    vd = shopify_graphql(verify_q, {"id": collection_id})["collection"]
    live_first_ids = [e["node"]["id"] for e in vd["products"]["edges"]]
    expected_first_ids = ordered_product_ids[:10]
    verified = live_first_ids == expected_first_ids

    return {
        "job_id": job_id, "job_done": done, "verified": verified,
        "live_sort_order": vd.get("sortOrder"),
        "live_first_titles": [e["node"]["title"] for e in vd["products"]["edges"]],
    }

def shopify_sync_collections():
    # Groups approved products by their own stored spreadsheet_type_id and syncs
    # each group to that type's configured `sync` collection list -- NOT the
    # generic `collection` field (used only for CSV Tags/Type). This is what
    # prevents the 2026-07-24 incident (2 Cut Beads getting blanket-synced into
    # "Czech Glass Beads" just because that was the category's tag value) from
    # happening again: a type with no `sync` entries is a safe no-op, reported
    # explicitly rather than silently doing nothing.
    products=load_products()
    results={"created_collections":[],"matched":0,"not_found_in_shopify":[],"skipped_no_target":[],"errors":[]}
    collection_cache={}
    by_type={}
    for p in products:
        if not p.get("approved") or p.get("skipped"): continue
        by_type.setdefault(p.get("spreadsheet_type_id") or "", []).append(p)
    for type_id, group in by_type.items():
        t = spreadsheet_type_by_id(type_id) if type_id else None
        sync_targets = (t or {}).get("sync", [])
        if not sync_targets:
            results["skipped_no_target"].extend({"handle":p["handle"],"type_id":type_id or "(none)"} for p in group)
            continue
        target_ids=[]
        for target in sync_targets:
            handle=target["handle"]
            if handle not in collection_cache:
                try:
                    cid, created=shopify_find_or_create_collection(target["title"], handle)
                    collection_cache[handle]=cid
                    if created: results["created_collections"].append(target["title"])
                except Exception as e:
                    results["errors"].append({"collection":target["title"],"error":str(e)}); continue
            target_ids.append(collection_cache[handle])
        for p in group:
            try:
                product_id=shopify_product_id_by_handle(p["handle"])
                if not product_id:
                    results["not_found_in_shopify"].append(p["handle"]); continue
                for cid in target_ids:
                    shopify_add_product_to_collection(product_id, cid)
                results["matched"]+=1
            except Exception as e:
                results["errors"].append({"handle":p["handle"],"error":str(e)})
    return results

# --- Neil's Color Family classification (see margola_color_family_taxonomy memory) ---
# Digit -> family map: this is Preciosa Ornela's own color-number convention
# (shared across all their bead lines), confirmed by Neil to be authoritative
# unless a listed exception below applies.
COLOR_FAMILY_DIGIT_MAP = {
    "0": "Crystal / White", "1": "Amber / Brown", "2": "Amethyst / Purple", "3": "Sapphire / Blue",
    "4": "Smoke Gray / Black Diamond", "5": "Green", "6": "Aqua / Turquoise", "7": "Pink",
    "8": "Yellow", "9": "Orange / Red",
}
# Individually confirmed overrides discovered auditing the Roller/Crow catalog --
# trusted to generalize to other categories since Preciosa's numbering is shared,
# but re-verify against a real product if a new category surfaces one of these
# codes with a name that doesn't match.
COLOR_FAMILY_FORCED = {
    "78102": "Crystal / White", "48102": "Crystal / White", "00050": "Crystal / White",
    "49102": "Smoke Gray / Black Diamond",  # "metallic" in the name, but confirmed to beat the metallic rule
    "23980": "Black",
    "14400": "Smoke Gray / Black Diamond",  # "GUNMETAL" -- not caught by the metallic-word rule below
    "59205": "Black",  # "JET BLACK IRIS" -- leading digit '5' says Green, Neil confirmed Black instead
    "5920": "Black",  # "JET BLACK ... MATTE" (code 5920M, strips to 5920) -- same override as 59205
}
# Only used as a fallback for names that don't follow the plain digit rule
# (composite crystal-lined codes, the 017xx series) -- NOT used to override or
# second-guess the digit rule for an ordinary code, since Neil confirmed the
# digit is authoritative even when it looks like it disagrees with the name
# (e.g. Teal correctly lands in Green, not Aqua/Turquoise, by digit).
COLOR_FAMILY_KEYWORDS = [
    ("TEAL", "Green"), ("PEACH", "Pink"), ("MAUVE", "Amethyst / Purple"), ("IVORY", "Crystal / White"),
    ("BLACK DIAMOND", "Smoke Gray / Black Diamond"), ("GRAY", "Smoke Gray / Black Diamond"), ("GREY", "Smoke Gray / Black Diamond"),
    ("SMOKE", "Smoke Gray / Black Diamond"), ("SAPPHIRE", "Sapphire / Blue"), ("BLUE", "Sapphire / Blue"),
    ("TURQUOISE", "Aqua / Turquoise"), ("AQUA", "Aqua / Turquoise"), ("EMERALD", "Green"), ("OLIVINE", "Green"),
    ("GREEN", "Green"), ("AMETHYST", "Amethyst / Purple"), ("VIOLET", "Amethyst / Purple"), ("PURPLE", "Amethyst / Purple"),
    ("FUCHSIA", "Pink"), ("ROSE", "Pink"), ("PINK", "Pink"), ("SIAM", "Orange / Red"), ("RUBY", "Orange / Red"),
    ("ORANGE", "Orange / Red"), ("RED", "Orange / Red"), ("CITRINE", "Yellow"), ("YELLOW", "Yellow"),
    ("ROOT BEER", "Amber / Brown"), ("BRONZE", "Amber / Brown"), ("AMBER", "Amber / Brown"), ("BROWN", "Amber / Brown"),
    ("BLACK", "Black"), ("WHITE", "Crystal / White"),
]
HORN_AGATE_CODE_RE = re.compile(r"^26[1-9]")
COLOR_FAMILY_CHOICES = ["Sapphire / Blue","Orange / Red","Green","Amber / Brown","Amethyst / Purple","Smoke Gray / Black Diamond","Aqua / Turquoise","Metallic","Black","Yellow","Crystal / White","Pink"]

def classify_color_family(color_number, color_name, description=""):
    # Returns (family_or_None, reason). family is None when no rule confidently
    # applies -- caller should flag it for a manual call rather than guess.
    code = re.sub(r"[^0-9]", "", clean(color_number))
    name = clean(color_name).upper()
    blob = f"{name} {re.sub('<[^>]+>',' ',clean(description)).upper()}"
    if code in COLOR_FAMILY_FORCED:
        return COLOR_FAMILY_FORCED[code], "forced exception"
    if "METALLIC" in blob:
        return "Metallic", "metallic in name/description"
    if HORN_AGATE_CODE_RE.match(code):
        return None, "Horn/Agate/Stone line uses a different numbering system"
    if code[:2] in ("37", "38"):
        lining = name.replace("CRYSTAL", "").replace("LUSTER", "")
        for kw, fam in COLOR_FAMILY_KEYWORDS:
            if kw in lining: return fam, f"crystal-lined composite, lining color matched '{kw}'"
        return None, "crystal-lined composite (37xxx/38xxx) -- lining color not recognized"
    if "IVORY" in blob:
        return "Crystal/White", "ivory"
    if code[:3] == "017":
        for kw, fam in COLOR_FAMILY_KEYWORDS:
            if kw in name: return fam, f"017xx excluded from digit rule, name matched '{kw}'"
        return None, "017xx series -- name not recognized"
    if code and code[0] in COLOR_FAMILY_DIGIT_MAP:
        return COLOR_FAMILY_DIGIT_MAP[code[0]], "leading digit"
    return None, "no recognizable color number"

def shopify_apply_color_family(items):
    # items: [{"handle":..., "family":...}, ...] -- confirmed by the user in the
    # review table, not computed fresh here, so an inline override in the UI is
    # respected exactly as typed.
    products_by_handle = {p["handle"]: p for p in load_products()}
    applied, errors = [], []
    for item in items:
        handle = clean(item.get("handle")); family = clean(item.get("family"))
        if not handle or not family:
            errors.append({"handle":handle,"error":"Missing handle or family"}); continue
        if family not in COLOR_FAMILY_CHOICES:
            errors.append({"handle":handle,"error":f"{family!r} is not one of the 12 approved families"}); continue
        p = products_by_handle.get(handle)
        try:
            product_id = shopify_product_id_by_handle(handle)
            if not product_id:
                errors.append({"handle":handle,"error":"Product not found on Shopify -- export/import it first"}); continue
            m="""mutation($metafields:[MetafieldsSetInput!]!){ metafieldsSet(metafields:$metafields){ metafields{id} userErrors{field message} } }"""
            mf=[{"ownerId":product_id,"namespace":"custom","key":"color_family","type":"single_line_text_field","value":family}]
            d=shopify_graphql(m,{"metafields":mf})["metafieldsSet"]
            if d.get("userErrors"): raise RuntimeError(json.dumps(d["userErrors"]))
            color_number = clean(p.get("color_number")) if p else ""
            if color_number:
                tm="""mutation($id:ID!,$tags:[String!]!){ tagsAdd(id:$id, tags:$tags){ userErrors{field message} } }"""
                td=shopify_graphql(tm,{"id":product_id,"tags":[color_number]})["tagsAdd"]
                if td.get("userErrors"): raise RuntimeError(json.dumps(td["userErrors"]))
            applied.append({"handle":handle,"family":family,"color_number":color_number})
        except Exception as e:
            errors.append({"handle":handle,"error":str(e)})
    return {"applied":applied,"errors":errors}

def taxonomy_map_path(): return os.path.join(DATA_DIR,"shopify_taxonomy_map.json")
def load_taxonomy_map():
    return json.load(open(taxonomy_map_path(),encoding="utf-8")) if os.path.exists(taxonomy_map_path()) else {"bead_shape":{},"color":{}}

COLOR_BASE_SWATCH_GIDS={
    "Beige":"gid://shopify/TaxonomyValue/6","Black":"gid://shopify/TaxonomyValue/1","Blue":"gid://shopify/TaxonomyValue/2",
    "Bronze":"gid://shopify/TaxonomyValue/657","Brown":"gid://shopify/TaxonomyValue/7","Clear":"gid://shopify/TaxonomyValue/17",
    "Gold":"gid://shopify/TaxonomyValue/4","Gray":"gid://shopify/TaxonomyValue/8","Green":"gid://shopify/TaxonomyValue/9",
    "Multicolor":"gid://shopify/TaxonomyValue/2865","Navy":"gid://shopify/TaxonomyValue/15","Orange":"gid://shopify/TaxonomyValue/10",
    "Pink":"gid://shopify/TaxonomyValue/11","Purple":"gid://shopify/TaxonomyValue/12","Red":"gid://shopify/TaxonomyValue/13",
    "Rose gold":"gid://shopify/TaxonomyValue/16","Silver":"gid://shopify/TaxonomyValue/5","White":"gid://shopify/TaxonomyValue/3",
    "Yellow":"gid://shopify/TaxonomyValue/14",
}
SOLID_PATTERN_GID="gid://shopify/TaxonomyValue/2874"
COLOR_FINISH_WORDS={"TRANSPARENT","TRANSPAENT","TRANSPAR","TRANS","PARENT","OPAQUE","MATTE"}
CODE_PREFIX_RE=re.compile(r"^[0-9A-Za-z]{4,6}\s*-\s*")

def normalize_color_key(label):
    s=CODE_PREFIX_RE.sub("", label or "")
    s=s.replace("&amp;","&").replace("&nbsp;"," ")
    s=re.sub(r"[^A-Za-z0-9]+"," ", s).upper().strip()
    return " ".join(w for w in s.split() if w not in COLOR_FINISH_WORDS)

def shopify_list_metaobjects(mo_type):
    nodes, cursor = [], None
    while True:
        q="""query($type:String!,$after:String){ metaobjects(type:$type, first:100, after:$after){ edges{ cursor node{ id handle displayName } } pageInfo{ hasNextPage } } }"""
        data=shopify_graphql(q,{"type":mo_type,"after":cursor})
        edges=data["metaobjects"]["edges"]
        nodes.extend(e["node"] for e in edges)
        if not data["metaobjects"]["pageInfo"]["hasNextPage"] or not edges: break
        cursor=edges[-1]["cursor"]
    return nodes

def shopify_create_color_pattern_metaobject(label, base_color_name):
    base_gid=COLOR_BASE_SWATCH_GIDS.get(base_color_name)
    if not base_gid: raise RuntimeError(f"No base swatch GID for {base_color_name!r}")
    m="""mutation($obj:MetaobjectCreateInput!){ metaobjectCreate(metaobject:$obj){ metaobject{ handle } userErrors{ field message } } }"""
    d=shopify_graphql(m,{"obj":{"type":"shopify--color-pattern","fields":[
        {"key":"label","value":label},
        {"key":"color_taxonomy_reference","value":json.dumps([base_gid])},
        {"key":"pattern_taxonomy_reference","value":SOLID_PATTERN_GID},
    ]}})["metaobjectCreate"]
    if d.get("userErrors"): raise RuntimeError(json.dumps(d["userErrors"]))
    return d["metaobject"]["handle"]

def shopify_create_size_metaobject(label, base_gid):
    m="""mutation($obj:MetaobjectCreateInput!){ metaobjectCreate(metaobject:$obj){ metaobject{ handle } userErrors{ field message } } }"""
    d=shopify_graphql(m,{"obj":{"type":"shopify--size","fields":[
        {"key":"label","value":label},
        {"key":"taxonomy_reference","value":base_gid},
    ]}})["metaobjectCreate"]
    if d.get("userErrors"): raise RuntimeError(json.dumps(d["userErrors"]))
    return d["metaobject"]["handle"]

class MetaobjectResolver:
    """Resolves Margola color/bead-shape text to Shopify metaobject handles for CSV export, creating new library entries when no match exists. Caches lookups for the lifetime of one export."""
    def __init__(self):
        self.taxonomy_map=load_taxonomy_map()
        self._color_by_key=None
        self._size_by_label=None
        self._created_colors={}
        self._created_sizes={}
        self.created_colors=[]
        self.created_sizes=[]
        self.unmapped_colors=[]
        self.unmapped_sizes=[]

    def _colors(self):
        if self._color_by_key is None:
            self._color_by_key={}
            for n in shopify_list_metaobjects("shopify--color-pattern"):
                self._color_by_key.setdefault(normalize_color_key(n["displayName"]), []).append(n)
        return self._color_by_key

    def _sizes(self):
        if self._size_by_label is None:
            self._size_by_label={n["displayName"].strip().upper():n for n in shopify_list_metaobjects("shopify--size")}
        return self._size_by_label

    def color_handle(self, color_name):
        color_name=clean(color_name)
        if not color_name: return ""
        key=normalize_color_key(color_name)
        by_key=self._colors()
        if key in by_key:
            candidates=by_key[key]
            plain=[c for c in candidates if not CODE_PREFIX_RE.match(c["displayName"])]
            return (plain[0] if plain else candidates[0])["handle"]
        if key in self._created_colors: return self._created_colors[key]
        base_color=self.taxonomy_map.get("color",{}).get(color_name)
        if not base_color:
            self.unmapped_colors.append(color_name); return ""
        label=key.title()
        handle=shopify_create_color_pattern_metaobject(label, base_color)
        self._created_colors[key]=handle
        self.created_colors.append(label)
        return handle

    def bead_shape_label(self, bead_shape):
        # Plain text, NOT json.dumps([...]) -- this feeds a CSV cell, and Shopify's
        # product CSV importer wraps a list.single_line_text_field cell's raw text
        # into a single-item list itself. Pre-JSON-encoding it here double-wraps
        # the value (confirmed live 2026-07-25: produced ["[\"Bugle\"]"] instead of
        # ["Bugle"]). json.dumps(...) is only correct for a direct GraphQL
        # metafieldsSet call (a different channel with different expectations),
        # not for a CSV import column -- don't copy formatting between the two
        # without verifying against that specific channel's actual behavior.
        spec=self.taxonomy_map.get("bead_shape",{}).get(clean(bead_shape))
        return spec["label"] if spec else ""

    def type_label(self, type_value):
        # Same plain-text-not-json.dumps rule as bead_shape_label() -- this feeds
        # the custom.type CSV column, and Shopify's CSV importer wraps the raw
        # cell text into a list itself for a list.single_line_text_field.
        return clean(type_value)

    def size_handle(self, size):
        spec=self.taxonomy_map.get("size",{}).get(clean(size))
        if not spec:
            if clean(size): self.unmapped_sizes.append(size)
            return ""
        label=spec["label"]; key=label.strip().upper()
        by_label=self._sizes()
        if key in by_label: return by_label[key]["handle"]
        if key in self._created_sizes: return self._created_sizes[key]
        handle=shopify_create_size_metaobject(label, spec["base_gid"])
        self._created_sizes[key]=handle
        self.created_sizes.append(label)
        return handle

class Handler(BaseHTTPRequestHandler):
    def send_json(self, obj, status=200):
        data = json.dumps(obj).encode("utf-8")
        self.send_response(status); self.send_header("Content-Type","application/json"); self.send_header("Content-Length",str(len(data))); self.end_headers(); self.wfile.write(data)
    def read_body(self): return self.rfile.read(int(self.headers.get("Content-Length","0")))
    def send_csv(self, text, filename):
        data = text.encode("utf-8-sig")
        self.send_response(200); self.send_header("Content-Type","text/csv; charset=utf-8"); self.send_header("Content-Disposition",f"attachment; filename={filename}"); self.send_header("Content-Length",str(len(data))); self.end_headers(); self.wfile.write(data)
    def serve_static(self,path):
        full = os.path.join(APP_DIR, path.lstrip("/"))
        if not os.path.abspath(full).startswith(APP_DIR) or not os.path.exists(full): self.send_response(404); self.end_headers(); return
        data = open(full,"rb").read()
        self.send_response(200); self.send_header("Content-Type",guess(path)); self.send_header("Content-Length",str(len(data))); self.end_headers(); self.wfile.write(data)
    def do_GET(self):
        parsed = urlparse(self.path)
        path = parsed.path
        query = parse_qs(parsed.query)
        if path == "/": return self.serve_static("/static/index.html")
        if path.startswith("/static/"): return self.serve_static(path)
        products = load_products()
        if path == "/api/products": return self.send_json({"summary":summary(products),"products":products})
        if path == "/api/export/review.csv": return self.send_csv(review_csv(products),"margola_review_report.csv")
        if path == "/api/export/shopify-preview.csv":
            try: return self.send_csv(shopify_csv(products, True),"margola_shopify_preview.csv")
            except Exception as e: return self.send_json({"ok":False,"error":str(e)},500)
        if path == "/api/export/shopify-test.csv":
            try: return self.send_csv(shopify_csv(products, True, 3),"margola_shopify_TEST_3_products.csv")
            except Exception as e: return self.send_json({"ok":False,"error":str(e)},500)
        if path == "/api/shopify/config":
            c=shopify_config(); return self.send_json({"store":c["store"],"has_client_id":bool(c["client_id"]),"has_client_secret":bool(c["client_secret"])})
        if path == "/api/shopify/test":
            try: return self.send_json({"ok":True,"shop":shopify_test_connection()})
            except Exception as e: return self.send_json({"ok":False,"error":str(e)},500)
        if path == "/api/shopify/files":
            try:
                files=shopify_list_files(); save_file_map(files); matched,total=apply_file_matches_to_products(); return self.send_json({"ok":True,"files":files,"matched":matched,"total_products":total})
            except Exception as e: return self.send_json({"ok":False,"error":str(e)},500)
        if path == "/api/spreadsheet-types":
            types=[{"id":t["id"],"label":t["label"],"sync":t.get("sync",[])} for t in load_spreadsheet_types()]
            return self.send_json({"ok":True,"types":types})
        if path == "/api/shopify/color-family-proposals":
            proposals=[]
            for p in products:
                if not p.get("approved") or p.get("skipped"): continue
                family, reason = classify_color_family(p.get("color_number"), p.get("color_name"), p.get("generated_description") or p.get("source_description"))
                proposals.append({"handle":p["handle"],"title":p.get("title"),"color_number":p.get("color_number"),
                                   "color_name":p.get("color_name"),"family":family,"reason":reason,"needs_review":family is None})
            return self.send_json({"ok":True,"proposals":proposals,"choices":COLOR_FAMILY_CHOICES})
        if path == "/api/shopify/seo-proposals":
            try: return self.send_json({"ok":True,"proposals":shopify_seo_proposals()})
            except Exception as e: return self.send_json({"ok":False,"error":str(e)},500)
        if path == "/api/shopify/collection-sort-proposal":
            handle = query.get("handle", [""])[0].strip()
            group_by_size = query.get("group_by_size", ["1"])[0] != "0"
            if not handle: return self.send_json({"ok":False,"error":"Missing handle"},400)
            try: return self.send_json({"ok":True, **shopify_collection_sort_proposal(handle, group_by_size)})
            except Exception as e: return self.send_json({"ok":False,"error":str(e)},500)
        if path == "/api/photos/resolve-from-drive":
            index = load_drive_index()
            if not index["found_any"]:
                return self.send_json({"ok":False,"error":f"No drive filelist found. Expected one of: {DRIVE_FILELISTS}"},400)
            proposals=[]
            for p in products:
                if p.get("image_src"): continue
                found = resolve_photo_from_drive(p, index)
                if found:
                    dims = _image_dimensions(found["source_path"]) if os.path.exists(found["source_path"]) else None
                    proposals.append({"id":p["id"],"title":p.get("title"),"factory_style":p.get("factory_style"),
                                       "source_path":found["source_path"],"target_filename":found["target_filename"],
                                       "reused_other_size":found["reused_other_size"],
                                       "dimensions": f"{dims[0]}x{dims[1]}" if dims else "unverified (drive not connected)",
                                       "correct_size": dims == TARGET_IMAGE_SIZE if dims else None})
            return self.send_json({"ok":True,"proposals":proposals,"filelists":DRIVE_FILELISTS})
        if path == "/api/photos/scan-mapped-folder":
            category = query.get("category", [""])[0]
            entry = PHOTO_MAPS.get(category)
            if not entry:
                return self.send_json({"ok":False,"error":f"Unknown category {category!r}. Known: {list(PHOTO_MAPS)}"},400)
            proposals=[]
            for p in products:
                if p.get("image_src"): continue
                filename = entry["map"].get(p.get("factory_style"))
                if not filename: continue
                source_path = f"{entry['folder']}/{filename}"
                dims = _image_dimensions(source_path) if os.path.exists(source_path) else None
                proposals.append({"id":p["id"],"title":p.get("title"),"factory_style":p.get("factory_style"),
                                   "source_path":source_path,"target_filename":clean(p.get("image_filename")),
                                   "reused_other_size":False,
                                   "dimensions": f"{dims[0]}x{dims[1]}" if dims else "unverified (drive not connected)",
                                   "correct_size": dims == TARGET_IMAGE_SIZE if dims else None})
            return self.send_json({"ok":True,"proposals":proposals,"folder":entry["folder"],"label":entry["label"]})
        self.send_response(404); self.end_headers()
    def do_POST(self):
        path = urlparse(self.path).path
        if path == "/api/shopify/sync-collections":
            try: return self.send_json({"ok":True, **shopify_sync_collections()})
            except Exception as e: return self.send_json({"ok":False,"error":str(e)},500)
        if path == "/api/shopify/apply-color-family":
            data = json.loads(self.read_body().decode("utf-8")); items = data.get("items",[])
            try: return self.send_json({"ok":True, **shopify_apply_color_family(items)})
            except Exception as e: return self.send_json({"ok":False,"error":str(e)},500)
        if path == "/api/shopify/apply-seo":
            data = json.loads(self.read_body().decode("utf-8")); items = data.get("items",[])
            try: return self.send_json({"ok":True, **shopify_apply_seo(items)})
            except Exception as e: return self.send_json({"ok":False,"error":str(e)},500)
        if path == "/api/shopify/apply-collection-sort":
            data = json.loads(self.read_body().decode("utf-8"))
            collection_id, product_ids = data.get("collection_id",""), data.get("product_ids",[])
            if not collection_id or not product_ids:
                return self.send_json({"ok":False,"error":"Missing collection_id or product_ids"},400)
            try: return self.send_json({"ok":True, **shopify_apply_collection_sort(collection_id, product_ids)})
            except Exception as e: return self.send_json({"ok":False,"error":str(e)},500)
        if path == "/api/photos/apply-drive-matches":
            data = json.loads(self.read_body().decode("utf-8")); matches = data.get("matches",[])
            files = load_file_map(); existing_by_filename = {f.get("filename"):f for f in files if f.get("filename")}
            uploaded=[]; errors=[]
            for m in matches:
                source_path, target_filename = m.get("source_path",""), m.get("target_filename","")
                if not source_path or not target_filename:
                    errors.append({"target_filename":target_filename,"error":"Missing source_path or target_filename"}); continue
                if target_filename in existing_by_filename:
                    uploaded.append({"target_filename":target_filename,"status":"already_known"}); continue
                if not os.path.exists(source_path):
                    errors.append({"target_filename":target_filename,"error":f"Drive not connected or file moved: {source_path}"}); continue
                try:
                    result = shopify_upload_file(source_path, alt=target_filename, filename=target_filename)
                    files.append(result); uploaded.append({"target_filename":target_filename,"url":result.get("url")})
                except Exception as e:
                    errors.append({"target_filename":target_filename,"error":str(e)})
            save_file_map(files); matched,total = apply_file_matches_to_products()
            return self.send_json({"ok":len(errors)==0,"uploaded_count":len(uploaded),"error_count":len(errors),
                                    "uploaded":uploaded,"errors":errors,"matched":matched,"total_products":total})
        if path == "/api/reset":
            save_products([]); return self.send_json({"ok":True,"summary":summary([])})
        if path == "/api/product/update":
            data = json.loads(self.read_body().decode("utf-8")); products = load_products()
            for p in products:
                if p["id"] == data.get("id"): p.update(data.get("updates",{})); break
            save_products(products); return self.send_json({"ok":True,"summary":summary(products)})
        if path == "/api/products/approve-with-image":
            products = load_products(); approved_count = 0
            for p in products:
                if p.get("image_src") and not p.get("skipped") and not p.get("approved"):
                    p["approved"] = True; p["status"] = "Approved"; approved_count += 1
            save_products(products)
            return self.send_json({"ok":True,"approved_count":approved_count,"summary":summary(products)})
        if path == "/api/descriptions/import":
            data = json.loads(self.read_body().decode("utf-8")); descs = parse_descriptions(data.get("text","")); products = load_products(); matched = 0
            for p in products:
                if p["handle"] in descs: p["generated_description"] = descs[p["handle"]]; p["description_approved"] = False; matched += 1
            save_products(products); return self.send_json({"ok":True,"matched":matched,"total_descriptions":len(descs),"summary":summary(products)})
        if path == "/api/shopify/upload-test-image":
            ctype=self.headers.get("Content-Type",""); m=re.search(r"boundary=(.+)", ctype)
            if not m: return self.send_json({"error":"Missing upload boundary"},400)

            body=self.read_body()
            boundary=("--"+m.group(1)).encode()
            uploaded=[]

            for part in body.split(boundary):
                if b'name="file"' not in part:
                    continue
                fm=re.search(rb'filename="([^"]+)"',part)
                if not fm:
                    continue
                filename=fm.group(1).decode("utf-8",errors="replace")
                if not filename:
                    continue
                pieces=part.split(b"\r\n\r\n",1)
                if len(pieces)!=2:
                    continue
                file_bytes=pieces[1].rstrip(b"\r\n--")
                if not file_bytes:
                    continue

                safe_name=os.path.basename(filename)
                upload_path=os.path.join(UPLOAD_DIR,"shopify_upload_"+safe_name)
                open(upload_path,"wb").write(file_bytes)
                uploaded.append((safe_name, upload_path))

            if not uploaded:
                return self.send_json({"error":"No images uploaded"},400)

            try:
                files=load_file_map()
                results=[]
                errors=[]

                existing_by_filename={f.get("filename"):f for f in files if f.get("filename")}

                for safe_name, upload_path in uploaded:
                    if safe_name in existing_by_filename:
                        results.append({"filename": safe_name, "status": "already_known", "url": existing_by_filename[safe_name].get("url","")})
                        continue

                    try:
                        result=shopify_upload_file(upload_path, alt=safe_name)
                        files.append(result)
                        results.append(result)
                    except Exception as item_error:
                        errors.append({"filename": safe_name, "error": str(item_error)})

                save_file_map(files)
                matched,total=apply_file_matches_to_products()

                return self.send_json({
                    "ok": len(errors)==0,
                    "uploaded_count": len(results),
                    "error_count": len(errors),
                    "files": results,
                    "errors": errors,
                    "matched": matched,
                    "total_products": total
                })
            except Exception as e:
                return self.send_json({"ok":False,"error":str(e)},500)
        if path == "/api/import":
            ctype = self.headers.get("Content-Type",""); m = re.search(r"boundary=(.+)", ctype)
            if not m: return self.send_json({"error":"Missing upload boundary"},400)
            body = self.read_body(); boundary = ("--"+m.group(1)).encode(); file_bytes = None; filename = "upload.xlsx"; type_id = None
            for part in body.split(boundary):
                if b'name="file"' in part:
                    fm = re.search(rb'filename="([^"]+)"', part)
                    if fm: filename = fm.group(1).decode("utf-8", errors="replace")
                    pieces = part.split(b"\r\n\r\n",1)
                    if len(pieces)==2: file_bytes = pieces[1].rstrip(b"\r\n--")
                elif b'name="type_id"' in part:
                    pieces = part.split(b"\r\n\r\n",1)
                    if len(pieces)==2: type_id = pieces[1].rstrip(b"\r\n--").decode("utf-8", errors="replace").strip()
            if not file_bytes: return self.send_json({"error":"No file uploaded"},400)
            upload_path = os.path.join(UPLOAD_DIR, f"{datetime.now().strftime('%Y%m%d_%H%M%S')}_{slugify(filename)}.xlsx")
            open(upload_path,"wb").write(file_bytes)
            try:
                products = parse_xlsx(upload_path, forced_type_id=(type_id or None)); save_products(products); return self.send_json({"summary":summary(products),"detected":detected_categories(products),"products":products})
            except Exception as e: return self.send_json({"error":str(e)},500)
        self.send_response(404); self.end_headers()

if __name__ == "__main__":
    port = int(os.environ.get("PORT","8787"))
    print(f"Margola Product Manager v3 running at http://localhost:{port}")
    ThreadingHTTPServer(("127.0.0.1", port), Handler).serve_forever()
